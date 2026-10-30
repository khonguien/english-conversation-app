#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generator script for:
1. materials/documents/speaking_test_questions.docx (Questions only, English only, numbered 1-29)
2. materials/documents/speaking_test_answers.md (Answers formatted turn-by-turn with clear spacing)
"""

import os
import re
import sys
import json
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

# Ensure utf-8 stdout
if sys.stdout.encoding != 'utf-8':
    sys.stdout.reconfigure(encoding='utf-8')

BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
DOCS_DIR = os.path.join(BASE_DIR, 'materials', 'documents')
TOPICS_DIR = os.path.join(BASE_DIR, 'src', 'data', 'topics')

# -------------------------------------------------------------
# 1. QUESTION PROMPTS (ENGLISH ONLY, NUMBERED 1 - 29)
# -------------------------------------------------------------

QUESTIONS = [
    # PART 1: AIRPORT (1 - 12)
    {
        "num": 1,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Check-in Counter",
        "prompt": "You are at the airport check-in counter for your flight to Bangkok. You want to present your passport, check in one piece of luggage, request an aisle seat, and ask whether you need to pick up your luggage during your layover in Singapore. How would you carry out this conversation with the check-in staff?"
    },
    {
        "num": 2,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Overweight Luggage",
        "prompt": "You are checking in your luggage at the airport counter, and the staff informs you that your bag is overweight by 3 kilograms. You want to ask how much the overweight fee is and decide whether to pay the fee or take some items out into your carry-on bag. How would you handle this conversation?"
    },
    {
        "num": 3,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Security Check",
        "prompt": "You are going through the airport security checkpoint. The security officer instructs you to place your personal items in the trays and asks about liquids and electronic devices. You want to clarify what items need to be removed from your bag. How would you interact with the security officer?"
    },
    {
        "num": 4,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Flight Delay",
        "prompt": "You are waiting at the departure gate and notice on the flight information screen that your flight to Bangkok has been delayed. You approach the airline ground staff to ask about the cause of the delay, the new boarding time, and whether food or drink vouchers are provided. What would you say?"
    },
    {
        "num": 5,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Finding Your Seat",
        "prompt": "You have boarded the airplane and are walking down the aisle looking for your seat (18F). You cannot find the row and your carry-on luggage does not fit easily into the overhead bin. How would you ask the flight attendant or a nearby passenger for assistance?"
    },
    {
        "num": 6,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Ordering Food & Drinks on the Plane",
        "prompt": "The flight attendants are serving meal options (chicken with rice or beef with noodles) and beverages in the cabin. You want to choose your meal, request a specific beverage with water, and ask for a cup of hot tea after your meal. How would you order from the flight attendant?"
    },
    {
        "num": 7,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Asking for Help on the Plane",
        "prompt": "You are on a long flight and the cabin temperature feels chilly. You want to request a blanket and an extra pillow, ask for a pen to fill in the arrival declaration form, and ask when the plane is expected to land. What would you say to the flight attendant?"
    },
    {
        "num": 8,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Immigration / Passport Control",
        "prompt": "You have just landed at an international airport and are standing before the Immigration Officer. The officer asks for your passport, the purpose of your visit, the duration of your stay, and where you will be staying. How would you answer clearly and present your travel documents?"
    },
    {
        "num": 9,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Customs Declaration",
        "prompt": "You have collected your luggage and are passing through the customs inspection lane. The customs officer asks if you have any items to declare, including fresh food, alcohol, cigarettes, or large amounts of cash. How would you respond to the officer?"
    },
    {
        "num": 10,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Baggage Claim & Information Desk",
        "prompt": "You arrive at the destination airport terminal. You want to ask the information desk staff where the baggage carousel for your flight is located, where you can exchange foreign currency, buy a local SIM card, and find a taxi stand to head into the city center. What would you say?"
    },
    {
        "num": 11,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Lost Luggage",
        "prompt": "You have waited at the baggage carousel for 30 minutes, but all luggage has cleared and your suitcase has not appeared. You walk to the Lost & Found / Baggage Service desk to report your missing bag, describe its appearance, and leave your contact and hotel information. How would you carry out this conversation?"
    },
    {
        "num": 12,
        "part": "PART 1: AT THE AIRPORT",
        "topic": "Airport",
        "title": "Asking for Directions in the Airport",
        "prompt": "You are in the airport departure hall and need to find the nearest restroom, the VAT refund counter, and the duty-free shopping area before heading to your gate. How would you ask the airport staff for directions?"
    },

    # PART 2: RESTAURANT (13 - 18)
    {
        "num": 13,
        "part": "PART 2: AT THE RESTAURANT",
        "topic": "Restaurant",
        "title": "Arrival, Reservation & Seating",
        "prompt": "You arrive at a restaurant for dinner. You want to state that you have a reservation for two at 7 PM under the name 'John' (or ask for an available table), request a quiet table by the window in the non-smoking section, and receive the menus. What would you say to the host?"
    },
    {
        "num": 14,
        "part": "PART 2: AT THE RESTAURANT",
        "topic": "Restaurant",
        "title": "Asking for Recommendations",
        "prompt": "You and your guest are seated at your table looking over the menu. You call the waiter over to ask for their top best-selling dishes, inquire about today's special, ask what the waiter personally recommends, and check whether a dish is gluten-free. How would you ask?"
    },
    {
        "num": 15,
        "part": "PART 2: AT THE RESTAURANT",
        "topic": "Restaurant",
        "title": "Full Course Ordering & Customizing",
        "prompt": "You are ready to order appetizers (chicken wings and French fries), main courses (ribeye steak cooked medium-rare with mashed potatoes; grilled chicken pasta less spicy and with no onions), and drinks (red wine and mineral water), requesting salad dressing on the side and extra cheese. How would you place this order with the waiter?"
    },
    {
        "num": 16,
        "part": "PART 2: AT THE RESTAURANT",
        "topic": "Restaurant",
        "title": "During the Meal & Handling Issues",
        "prompt": "During your meal, you ask the waiter for extra cutlery, fresh napkins, water, and warm bread. Later, you notice that a dish served is not what you ordered and the food is cold. How do you politely call the server over and ask for a fresh replacement?"
    },
    {
        "num": 17,
        "part": "PART 2: AT THE RESTAURANT",
        "topic": "Restaurant",
        "title": "Food Review & Dessert",
        "prompt": "The waiter approaches your table to ask how everything is tasting. You want to compliment the delicious meal, tell the server that you are quite full, but ask them to surprise you with the house's best signature dessert to share. What would you say?"
    },
    {
        "num": 18,
        "part": "PART 2: AT THE RESTAURANT",
        "topic": "Restaurant",
        "title": "Bill, Leftovers To-Go & Payment",
        "prompt": "You have finished dining. You want to ask for a to-go box for leftover food, request the bill, ask to split the total evenly between two cards, check whether the service charge is already included, and leave a generous tip for the wonderful service. How would you carry out this conversation?"
    },

    # PART 3: HOTEL (19 - 22)
    {
        "num": 19,
        "part": "PART 3: AT THE HOTEL",
        "topic": "Hotel",
        "title": "Check-in & Sea View Room",
        "prompt": "You arrive at the hotel reception desk to check in. You present your reservation under your name, show your ID, ask if an upgrade to a room with a sea view is available, and ask about breakfast serving hours and the Wi-Fi password. How would you talk to the receptionist?"
    },
    {
        "num": 20,
        "part": "PART 3: AT THE HOTEL",
        "topic": "Hotel",
        "title": "Housekeeping & Room Service",
        "prompt": "You are in your hotel room. The housekeeping staff knocks on your door to clean the room, change bedsheets, and provide fresh towels and bottled water. Later in the evening, you call room service to order dinner delivered to your room. What would you say in these interactions?"
    },
    {
        "num": 21,
        "part": "PART 3: AT THE HOTEL",
        "topic": "Hotel",
        "title": "Bathroom Plumbing Issues",
        "prompt": "You discover severe issues in your hotel bathroom: the toilet is clogged, the sink tap is leaking, and there is no hot water in the shower. You call the front desk to report the problems and urgently request maintenance or a room transfer. What would you say?"
    },
    {
        "num": 22,
        "part": "PART 3: AT THE HOTEL",
        "topic": "Hotel",
        "title": "Local Recommendations & Check-out",
        "prompt": "You ask the hotel receptionist for directions to the nearest beach and a recommendation for a good local seafood restaurant. Later on, you return to the front desk to check out, settle your incidental bill with a credit card, and request a receipt. How would you handle this conversation?"
    },

    # PART 4: CAFE (23 - 25)
    {
        "num": 23,
        "part": "PART 4: AT THE CAFE",
        "topic": "Cafe",
        "title": "Ordering at the Counter & Customizing",
        "prompt": "You walk up to the counter at a busy coffee shop. You ask the barista for the shop's best-selling drink, order a large iced coffee with milk, customize it with less ice and no sugar, specify that it is to go, and politely decline an extra pastry. How would you carry out this conversation?"
    },
    {
        "num": 24,
        "part": "PART 4: AT THE CAFE",
        "topic": "Cafe",
        "title": "Toppings, QR Payment & Order Status",
        "prompt": "You are at the counter ordering drinks for yourself and a colleague. You want to order two specialty drinks with extra boba and cheese foam, pay via bank transfer by scanning the counter QR code, and ask where to wait to pick up your drinks. What would you say to the barista?"
    },
    {
        "num": 25,
        "part": "PART 4: AT THE CAFE",
        "topic": "Cafe",
        "title": "Pick-up, Seating, Wi-Fi & Cafe Amenities",
        "prompt": "You pick up your drinks at the counter and ask for straws and extra ice. You ask whether there is available seating upstairs, inquire about the Wi-Fi password, and ask the staff to wipe down a dirty table and adjust the cold air conditioner above your seat. What would you say?"
    },

    # PART 5: WORKPLACE (26 - 29)
    {
        "num": 26,
        "part": "PART 5: AT THE WORKPLACE",
        "topic": "Workplace",
        "title": "What Do You Do for a Living?",
        "prompt": "You meet a friend or acquaintance after some time apart. You ask what each other does for a living, talk about the differences between remote freelance work and working in an office at a logistics company, and discuss the daily work atmosphere. How would you carry out this conversation?"
    },
    {
        "num": 27,
        "part": "PART 5: AT THE WORKPLACE",
        "topic": "Workplace",
        "title": "Daily Challenges & Deadlines",
        "prompt": "During lunch or an after-work coffee break, you and your coworker talk about the greatest pressures of your job, dealing with strict shipping schedules and tight closing deadlines, working overtime, and trying to maintain a healthy work-life balance. How would you discuss this?"
    },
    {
        "num": 28,
        "part": "PART 5: AT THE WORKPLACE",
        "topic": "Workplace",
        "title": "Five-Year Vision & Promotion Banter",
        "prompt": "You and a close coworker are relaxing after work and playfully banter about where you see yourselves in 5 years, teasing each other about seeking promotions, becoming managers, dealing with corporate stress, and negotiating higher pay. How would you carry out this playful conversation?"
    },
    {
        "num": 29,
        "part": "PART 5: AT THE WORKPLACE",
        "topic": "Workplace",
        "title": "Taking the Leap & Resignation",
        "prompt": "During an office break, you and a colleague discuss the news that a long-time teammate has recently submitted their resignation letter to leave the corporate job and start a new business venture. How do you discuss their decision and wish them future success?"
    }
]

# -------------------------------------------------------------
# 2. RESTAURANT VIETNAMESE TRANSLATIONS
# -------------------------------------------------------------

RESTAURANT_TURNS_VI = [
    # Scenario 1 (10 turns)
    [
        "Xin chào, tôi có thể giúp gì cho các bạn ạ? Nhóm mình đi mấy người thế ạ?",
        "Xin chào, cho tôi bàn 2 người nhé. Hiện tại quán còn bàn trống không?",
        "Quý khách đã đặt bàn trước với chúng tôi chưa ạ?",
        "Vâng, tôi có đặt bàn 2 người lúc 7 giờ tối dưới tên John.",
        "Để tôi kiểm tra ngay ạ. Vâng, anh John, bàn của anh đã sẵn sàng rồi! Mời đi lối này ạ.",
        "Cho chúng tôi xin bàn gần cửa sổ được không? Chúng tôi thích một bàn yên tĩnh nếu có thể.",
        "Để tôi kiểm tra xem bàn đó có ai đặt chưa; nếu chưa, mời quý khách qua đó. Vâng, bàn trống ạ! Ngay đây.",
        "Bàn này tuyệt quá, cảm ơn bạn. Và khu vực không hút thuốc nhé.",
        "Tất cả chỗ ngồi trong nhà của chúng tôi đều không hút thuốc ạ. Đây là thực đơn. Tôi sẽ để quý khách vài phút để xem qua.",
        "Thế này là hoàn hảo rồi, cảm ơn bạn."
    ],
    # Scenario 2 (9 turns)
    [
        "Xin chào, hôm nay tôi có thể phục vụ gì cho quý khách ạ? Quý khách đã sẵn sàng gọi món chưa?",
        "Xin lỗi, chúng tôi gọi món bây giờ được không? Nhưng trước tiên, bạn gợi ý món gì nào? Món nào bán chạy nhất ở đây?",
        "Bít tết ribeye và mỳ Ý hải sản chắc chắn là những món bán chạy nhất của chúng tôi ạ.",
        "Món đặc biệt trong ngày hôm nay là gì? Cá nhân bạn thì sẽ chọn món nào?",
        "Món đặc biệt hôm nay là cá hồi nướng ăn kèm măng tây ạ. Cá nhân tôi thì đó là món tôi cực kỳ yêu thích.",
        "Nghe tuyệt quá! Món đó có không chứa gluten (gluten-free) không bạn?",
        "Dạ vâng, món đó hoàn toàn không chứa gluten ạ.",
        "Tuyệt vời. Chúng tôi xin một phút để chọn món chính nhé.",
        "Quý khách cứ thong thả ạ, chúng tôi không vội đâu."
    ],
    # Scenario 3 (13 turns)
    [
        "Chào quý khách! Quý khách đã sẵn sàng gọi món chưa ạ?",
        "Vâng, chúng tôi sẵn sàng rồi. Về món khai vị, cho tôi cánh gà cùng với khoai tây chiên nhé.",
        "Lựa chọn tuyệt vời ạ. Thế còn món chính quý khách dùng gì ạ?",
        "Cho tôi bít tết ribeye kèm khoai tây nghiền nhé.",
        "Lựa chọn tuyệt hảo! Quý khách muốn bít tết nướng ở độ chín nào ạ?",
        "Cho tôi mức tái vừa (medium-rare) nhé.",
        "Đã ghi nhận: bít tết ribeye tái vừa kèm khoai tây nghiền. Thế còn người bạn đi cùng thì sao ạ?",
        "Tôi muốn mỳ Ý gà nướng, nhưng làm ít cay giúp tôi nhé? Và đừng bỏ hành tây.",
        "Đã ghi chú: mỳ Ý gà nướng, ít cay, không hành tây. Quý khách có muốn sốt để riêng một bên không ạ?",
        "Vâng, thêm phô mai và sốt để riêng giúp tôi nhé.",
        "Vâng ạ. Thế quý khách muốn dùng đồ uống gì ạ?",
        "Cho tôi một ly rượu vang đỏ và một chai nước khoáng nhé?",
        "Dạ được chứ ạ! Tôi sẽ quay lại ngay với đồ uống và món khai vị cho quý khách."
    ],
    # Scenario 4 (8 turns)
    [
        "Em ơi, cho chúng tôi xin thêm dao nĩa được không? Và cho tôi thêm một khăn ăn nữa nhé?",
        "Dạ có ngay ạ, đây là bộ dao nĩa mới và khăn ăn cho quý khách.",
        "Cho chúng tôi thêm ít nước lọc và thêm bánh mì nữa được không?",
        "Dạ chắc chắn rồi ạ, tôi sẽ nhờ bếp chuẩn bị thêm một giỏ bánh mì nóng cho quý khách.",
        "Xin lỗi bạn, tôi nghĩ món này không phải món tôi gọi. Hình như có sự nhầm lẫn với đơn của tôi, và thức ăn cũng bị nguội nữa.",
        "Tôi vô cùng xin lỗi về điều đó! Để tôi kiểm tra lại hóa đơn ngay ạ. Thành thật xin lỗi quý khách vì sự bất tiện này.",
        "Bạn có thể đổi cho tôi một đĩa mới nóng hổi được không?",
        "Dạ chắc chắn rồi ạ! Chúng tôi sẽ nhờ bếp làm lại món mới ngay lập tức. Chúng tôi sẽ mang ra sớm nhất có thể ạ."
    ],
    # Scenario 5 (7 turns)
    [
        "Chào quý khách, mọi thứ thế nào ạ? Bữa ăn dùng có vừa miệng không ạ?",
        "Thức ăn xuất sắc lắm! Miếng bít tết thật sự rất ngon và nướng chuẩn độ chín. Cảm ơn bạn rất nhiều.",
        "Tôi rất vui khi nghe điều đó! Thật tuyệt. Quý khách có muốn gọi thêm món tráng miệng nào không ạ?",
        "Dạ thôi, tôi no căng bụng rồi, cảm ơn bạn!",
        "Quý khách có muốn dùng cà phê hay trà, hoặc chọn một món tráng miệng nhẹ nhàng không ạ?",
        "Thực ra thì bạn cứ mang ra món tráng miệng ngon nhất của quán để chúng tôi cùng thưởng thức nhé.",
        "Dạ có ngay ạ! Tôi sẽ mang món tráng miệng đặc sản của quán ra ngay."
    ],
    # Scenario 6 (13 turns)
    [
        "Tôi có thể giúp thêm gì cho quý khách nữa không ạ?",
        "Không, chúng tôi ổn rồi, cảm ơn bạn. Cho tôi xin hóa đơn và một hộp mang về để gói phần thức ăn thừa này nhé?",
        "Dạ tất nhiên rồi ạ. Tôi xin để hóa đơn ở đây, và tôi sẽ quay lại ngay với hộp mang về ạ.",
        "Bạn có thể chia đôi hóa đơn cho chúng tôi được không?",
        "Dạ được chứ ạ, quý khách muốn chia đều 50-50 đúng không ạ?",
        "Đúng rồi bạn. Tôi thanh toán bằng thẻ được chứ?",
        "Dạ vâng, bên tôi chấp nhận tất cả các loại thẻ chính. Quý khách có thể chạm hoặc cắm thẻ ngay đây ạ.",
        "Phí phục vụ đã bao gồm trong tổng tiền chưa bạn?",
        "Dạ vâng, phí phục vụ đã được tính gộp trong tổng số tiền rồi ạ.",
        "Tuyệt quá. Bạn cứ giữ tiền thối làm tiền tip cho dịch vụ tuyệt vời hôm nay nhé!",
        "Quý khách thật tốt bụng, cảm ơn rất nhiều ạ! Cảm ơn quý khách đã dùng bữa tại nhà hàng hôm nay!",
        "Cảm ơn bạn, bữa ăn rất tuyệt vời! Chúc bạn một ngày tốt lành!",
        "Cảm ơn quý khách! Chúc quý khách một ngày tuyệt vời phía trước, mong sớm gặp lại quý khách ạ!"
    ]
]

# -------------------------------------------------------------
# 3. DIALOGUE TURN FORMATTING (TURN-BY-TURN CARD)
# -------------------------------------------------------------

def get_role_avatar(role_name):
    r = role_name.lower()
    if any(k in r for k in ['staff', 'receptionist', 'front desk']):
        return "🛎️"
    if any(k in r for k in ['waiter', 'server', 'bồi bàn']):
        return "🛎️"
    if any(k in r for k in ['officer']):
        return "👮"
    if any(k in r for k in ['flight attendant', 'tiếp viên']):
        return "✈️"
    if any(k in r for k in ['barista']):
        return "☕"
    if any(k in r for k in ['housekeeping']):
        return "🧹"
    if any(k in r for k in ['room service']):
        return "🍽️"
    if any(k in r for k in ['friend a', 'colleague a']):
        return "👥"
    if any(k in r for k in ['friend b', 'colleague b']):
        return "🧑‍💼"
    return "👤"

def format_turn(turn):
    role = turn.get('roleName', 'Speaker')
    text_en = turn.get('textEn', '').strip()
    text_vi = turn.get('textVi', '').strip()
    avatar = get_role_avatar(role)

    # Action narration e.g. *(Passenger takes some items out)*
    if text_en.startswith('*(') or text_en.startswith('('):
        clean_en = text_en.strip('*() ')
        clean_vi = text_vi.strip('*() ') if text_vi else ''
        res = f"📦 *({clean_en})*  "
        if clean_vi:
            res += f"\n*({clean_vi})*"
        return res

    # Spoken dialogue turn: English bold role + quotes, then soft break (2 spaces), then Vietnamese in italics
    lines = [f"{avatar} **{role}:** \"{text_en}\"  "]
    if text_vi:
        lines.append(f"*({text_vi})*")
    return "\n".join(lines)

# -------------------------------------------------------------
# 4. LOAD ALL 29 SCENARIOS FROM JSON + TRANSLATIONS
# -------------------------------------------------------------

def load_all_scenarios():
    scenarios = []

    # 1. Airport (12)
    with open(os.path.join(TOPICS_DIR, 'airport.json'), 'r', encoding='utf-8') as f:
        a_data = json.load(f)
    for s in a_data['scenarios']:
        if 'freetalk' in s['id']:
            continue
        scenarios.append(s['sampleDialogue'])

    # 2. Restaurant (6)
    with open(os.path.join(TOPICS_DIR, 'restaurant.json'), 'r', encoding='utf-8') as f:
        r_data = json.load(f)
    r_idx = 0
    for s in r_data['scenarios']:
        if 'freetalk' in s['id']:
            continue
        turns = s['sampleDialogue']
        vi_list = RESTAURANT_TURNS_VI[r_idx]
        for t_idx, turn in enumerate(turns):
            if t_idx < len(vi_list):
                turn['textVi'] = vi_list[t_idx]
        scenarios.append(turns)
        r_idx += 1

    # 3. Hotel (4)
    with open(os.path.join(TOPICS_DIR, 'hotel.json'), 'r', encoding='utf-8') as f:
        h_data = json.load(f)
    for s in h_data['scenarios']:
        if 'freetalk' in s['id']:
            continue
        scenarios.append(s['sampleDialogue'])

    # 4. Cafe (3)
    with open(os.path.join(TOPICS_DIR, 'cafe.json'), 'r', encoding='utf-8') as f:
        c_data = json.load(f)
    for s in c_data['scenarios']:
        if 'freetalk' in s['id']:
            continue
        scenarios.append(s['sampleDialogue'])

    # 5. Workplace (4)
    with open(os.path.join(TOPICS_DIR, 'workplace.json'), 'r', encoding='utf-8') as f:
        w_data = json.load(f)
    for s in w_data['scenarios']:
        if 'freetalk' in s['id']:
            continue
        scenarios.append(s['sampleDialogue'])

    return scenarios

# -------------------------------------------------------------
# 5. BUILD DOCX (QUESTIONS ONLY, ENGLISH ONLY, SIMPLE & CLEAN)
# -------------------------------------------------------------

def build_docx_questions(docx_path):
    doc = Document()
    
    # Page setup: Standard Margins 1 inch
    for section in doc.sections:
        section.top_margin = Inches(1.0)
        section.bottom_margin = Inches(1.0)
        section.left_margin = Inches(1.0)
        section.right_margin = Inches(1.0)
    
    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_title = title_p.add_run("ENGLISH SPEAKING TEST BANK")
    r_title.bold = True
    r_title.font.name = "Calibri"
    r_title.font.size = Pt(20)
    r_title.font.color.rgb = RGBColor(31, 78, 120)

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_before = Pt(0)
    sub_p.paragraph_format.space_after = Pt(18)
    sub_p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_sub = sub_p.add_run("29 Real-World Situation Prompts (Airport • Restaurant • Hotel • Cafe • Workplace)")
    r_sub.font.name = "Calibri"
    r_sub.font.size = Pt(11)
    r_sub.font.color.rgb = RGBColor(89, 89, 89)

    current_part = None
    for q in QUESTIONS:
        if q["part"] != current_part:
            current_part = q["part"]
            part_p = doc.add_paragraph()
            part_p.paragraph_format.space_before = Pt(16)
            part_p.paragraph_format.space_after = Pt(8)
            part_p.paragraph_format.keep_with_next = True
            r_part = part_p.add_run(current_part)
            r_part.bold = True
            r_part.font.name = "Calibri"
            r_part.font.size = Pt(13)
            r_part.font.color.rgb = RGBColor(31, 78, 120)

        # Question heading
        q_heading = doc.add_paragraph()
        q_heading.paragraph_format.space_before = Pt(8)
        q_heading.paragraph_format.space_after = Pt(2)
        q_heading.paragraph_format.keep_with_next = True
        r_qhead = q_heading.add_run(f"Question {q['num']}: {q['title']}")
        r_qhead.bold = True
        r_qhead.font.name = "Calibri"
        r_qhead.font.size = Pt(11.5)
        r_qhead.font.color.rgb = RGBColor(0, 0, 0)

        # Question prompt
        prompt_p = doc.add_paragraph()
        prompt_p.paragraph_format.space_before = Pt(0)
        prompt_p.paragraph_format.space_after = Pt(8)
        prompt_p.paragraph_format.line_spacing = 1.15
        r_prompt = prompt_p.add_run(q["prompt"])
        r_prompt.font.name = "Calibri"
        r_prompt.font.size = Pt(11)
        r_prompt.font.color.rgb = RGBColor(38, 38, 38)

    doc.save(docx_path)
    print(f"Saved questions Word document to: {docx_path}")

# -------------------------------------------------------------
# 6. BUILD MARKDOWN (ANSWERS ONLY, CLEAR TURN-BY-TURN FORMATTING)
# -------------------------------------------------------------

def build_md_answers(md_path):
    all_scenarios = load_all_scenarios()

    if len(all_scenarios) != len(QUESTIONS):
        print(f"WARNING: Loaded {len(all_scenarios)} scenarios but have {len(QUESTIONS)} questions!")

    lines = []
    lines.append("# 🎙️ BỘ ĐÁP ÁN ĐỀ THI NÓI TIẾNG ANH THỰC CHIẾN (29 TÌNH HUỐNG)")
    lines.append("### (SPEAKING TEST ANSWER KEY: 29 SAMPLE ROLEPLAY DIALOGUES)")
    lines.append("")
    lines.append("> **Tài liệu tham chiếu:** Tương ứng 1-1 với 29 câu hỏi trong file đề bài [`speaking_test_questions.docx`](file:///D:/Desktop/English%20in%20Conservation/materials/documents/speaking_test_questions.docx).  ")
    lines.append("> **Quy cách trình bày:** Mỗi câu thoại là một lượt đối đáp độc lập, có biểu tượng avatar, tên vai diễn in đậm, câu tiếng Anh nổi bật kèm bản dịch tiếng Việt ngay bên dưới.")
    lines.append("")
    lines.append("---")
    lines.append("")
    lines.append("## 📌 MỤC LỤC ĐÁP ÁN (TABLE OF ANSWERS)")
    lines.append("")

    current_part = None
    for q in QUESTIONS:
        if q["part"] != current_part:
            current_part = q["part"]
            lines.append(f"### {current_part}")
        clean_title = re.sub(r'[^a-zA-Z0-9\s-]', '', f"answer {q['num']} {q['title']}").strip().lower()
        anchor = re.sub(r'\s+', '-', clean_title)
        lines.append(f"- [Answer {q['num']}: {q['title']}](#{anchor})")
    lines.append("")
    lines.append("---")
    lines.append("")

    current_part = None
    for i, q in enumerate(QUESTIONS):
        if q["part"] != current_part:
            current_part = q["part"]
            lines.append(f"# {current_part}")
            lines.append("")

        turns = all_scenarios[i] if i < len(all_scenarios) else []

        lines.append(f"## Answer {q['num']}: {q['title']}")
        lines.append("")
        lines.append(f"> **Reference Question:** Question {q['num']} ({q['topic']})  ")
        lines.append(f"> **Question Prompt:** *\"{q['prompt']}\"*")
        lines.append("")
        lines.append("#### 💬 Kịch bản hội thoại mẫu (Sample Dialogue):")
        lines.append("")

        # Format each turn cleanly with an empty line between turns
        formatted_turns = [format_turn(t) for t in turns]
        lines.append("\n\n".join(formatted_turns))
        lines.append("")
        lines.append("---")
        lines.append("")

    with open(md_path, 'w', encoding='utf-8') as f:
        f.write('\n'.join(lines))
    print(f"Saved formatted answers Markdown document to: {md_path}")

# -------------------------------------------------------------
# MAIN
# -------------------------------------------------------------

def main():
    docx_file = os.path.join(DOCS_DIR, 'speaking_test_questions.docx')
    md_file = os.path.join(DOCS_DIR, 'speaking_test_answers.md')

    print("Step 1: Building Word questions document...")
    build_docx_questions(docx_file)

    print("Step 2: Building Markdown answers document (reformatted turn-by-turn)...")
    build_md_answers(md_file)

    print("Done! Both files generated successfully.")

if __name__ == '__main__':
    main()
