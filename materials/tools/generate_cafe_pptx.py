# -*- coding: utf-8 -*-
"""Generate Cafe.pptx presentation using Airport_backup.pptx template for consistent styling."""

import os
import sys
from pptx import Presentation
from pptx.util import Pt
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding='utf-8')

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
PROJECT_DIR = os.path.abspath(os.path.join(SCRIPT_DIR, "..", ".."))
TEMPLATE_PATH = os.path.join(PROJECT_DIR, "materials", "slides", "Airport_backup.pptx")
OUTPUT_PATH = os.path.join(PROJECT_DIR, "materials", "slides", "Cafe.pptx")

def clean_presentation(prs):
    while len(prs.slides) > 0:
        sldId = prs.element.sldIdLst[0]
        rId = sldId.rId
        prs.part.drop_rel(rId)
        prs.element.sldIdLst.remove(sldId)

def add_p(tf, text, is_bold=False, is_italic=False, font_size=None):
    if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == '':
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.text = text
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', '0')
    pPr.set('indent', '0')
    ns = nsdecls('a')
    buNone = parse_xml(f'<a:buNone {ns}/>')
    pPr.append(buNone)
    for r in p.runs:
        if is_bold:
            r.font.bold = True
        if is_italic:
            r.font.italic = True
        if font_size:
            r.font.size = Pt(font_size)
    return p

def add_vocab_slide(prs, category, word, example):
    slide = prs.slides.add_slide(prs.slide_layouts[1]) # Title and Content
    
    # Title
    title = slide.shapes.title
    title.text = "Vocabulary"
    for r in title.text_frame.paragraphs[0].runs:
        r.font.bold = True
    
    # Body
    body = slide.shapes.placeholders[1]
    body.left = 838200
    body.top = 1265125
    body.width = 10515600
    body.height = 4911838
    
    tf = body.text_frame
    tf.word_wrap = True
    
    add_p(tf, category, is_bold=True, font_size=20)
    add_p(tf, "")
    add_p(tf, word, is_bold=False, font_size=18)
    add_p(tf, example, is_italic=True, font_size=16)
    return slide

def add_content_slide(prs, title_text, subtitle_text, lines, font_size=13):
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = title_text
    for r in slide.shapes.title.text_frame.paragraphs[0].runs:
        r.font.bold = True
    
    body = slide.shapes.placeholders[1]
    body.left = 838200
    body.top = 1265125
    body.width = 10515600
    body.height = 4911838
    
    tf = body.text_frame
    tf.word_wrap = True
    
    if subtitle_text:
        add_p(tf, subtitle_text, is_bold=True, font_size=15)
        add_p(tf, "")
        
    for line in lines:
        if line.startswith("- ") or line.startswith("• "):
            add_p(tf, line, font_size=font_size)
        elif ":" in line and not line.startswith(" "):
            add_p(tf, line, is_bold=True, font_size=font_size)
        elif line == "":
            add_p(tf, "", font_size=font_size // 2)
        else:
            add_p(tf, line, font_size=font_size)
    return slide

def main():
    print(f"Loading template: {TEMPLATE_PATH}")
    prs = Presentation(TEMPLATE_PATH)
    clean_presentation(prs)
    print("Cleaned existing slides. Generating Cafe slides...")

    # Slide 1: Title
    s1 = prs.slides.add_slide(prs.slide_layouts[0])
    s1.shapes.title.text = "At the Cafe & Coffee Shop ☕"
    for r in s1.shapes.title.text_frame.paragraphs[0].runs:
        r.font.bold = True
    s1.shapes.placeholders[1].text = "English in Conversation — Tiếng Anh Giao Tiếp Quán Cà Phê\nTừ vựng trọng tâm • 3 Đoạn hội thoại tự nhiên • 43 Mẫu câu phản xạ"

    # Slide 2: Agenda
    agenda_lines = [
        "PHẦN 1: TỪ VỰNG TRỌNG TÂM (CORE VOCABULARY)",
        "• Đồ uống & Tùy chỉnh: iced coffee with milk, best-selling drink, sweetness level, pastry",
        "• Topping & Thanh toán: topping, boba, cheese foam, bank transfer, receipt",
        "• Tiện ích quán: straw, upstairs, Wi-Fi password, air conditioner, wipe, napkin",
        "",
        "PHẦN 2: 3 ĐOẠN HỘI THOẠI TỰ NHIÊN (REAL CONVERSATIONS)",
        "• Hội thoại 1: Gọi món tại quầy & Tùy chỉnh khẩu vị (Ordering & Customizing)",
        "• Hội thoại 2: Thêm Topping, Thanh toán QR & Chờ lấy đồ (Toppings, QR & Waiting)",
        "• Hội thoại 3: Nhận đồ, Tiện ích chỗ ngồi & Chăm sóc khách (Pick-up & Amenities)",
        "",
        "PHẦN 3: 43 MẪU CÂU GIAO TIẾP PHẢN XẠ THỰC CHIẾN (ESSENTIAL PHRASES)",
        "• 21 Mẫu câu của Nhân viên quầy & Phục vụ (Staff Phrases)",
        "• 22 Mẫu câu của Khách hàng (Customer Phrases)"
    ]
    add_content_slide(prs, "LESSON OUTLINE / MỤC LỤC BÀI HỌC", "", agenda_lines, font_size=14)

    # Core Vocabulary Slides
    vocab_items = [
        {
            "cat": "TỪ VỰNG: ĐỒ UỐNG & GỌI MÓN (1/3)",
            "word": "• iced coffee with milk /ˌaɪst ˈkɔːfi wɪð mɪlk/ (n phr): cà phê sữa đá",
            "ex": "  Ví dụ: Can I get an iced coffee with milk, please? (Cho em một ly cà phê sữa đá.)"
        },
        {
            "cat": "TỪ VỰNG: ĐỒ UỐNG & GỌI MÓN (1/3)",
            "word": "• best-selling drink /ˌbestˈselɪŋ drɪŋk/ (n phr): đồ uống bán chạy nhất",
            "ex": "  Ví dụ: What's your best-selling drink here? (Món nào hot nhất ở quán mình đấy chị?)"
        },
        {
            "cat": "TỪ VỰNG: ĐỒ UỐNG & GỌI MÓN (1/3)",
            "word": "• for here or to go /fər hɪr ɔːr tə ɡoʊ/ (phrase): dùng tại quán hay mang đi",
            "ex": "  Ví dụ: Will that latte be for here or to go? (Đồ uống đó mình dùng tại quán hay mang đi ạ?)"
        },
        {
            "cat": "TỪ VỰNG: ĐỒ UỐNG & GỌI MÓN (1/3)",
            "word": "• pastry /ˈpeɪstri/ (n): bánh ngọt / bánh nướng ăn kèm",
            "ex": "  Ví dụ: Would you like a pastry with that? (Mình có muốn thêm bánh không ạ?)"
        },
        {
            "cat": "TỪ VỰNG: TOPPING & THANH TOÁN QR (2/3)",
            "word": "• topping /ˈtɑːpɪŋ/ (n): đồ ăn kèm thêm trên đồ uống (trân châu, thạch, kem)",
            "ex": "  Ví dụ: What toppings do you have? (Có topping gì đấy ạ?)"
        },
        {
            "cat": "TỪ VỰNG: TOPPING & THANH TOÁN QR (2/3)",
            "word": "• boba /ˈboʊbə/ (n): trân châu (làm từ bột sắn)",
            "ex": "  Ví dụ: Can I add boba? (Cho em thêm trân châu nhá.)"
        },
        {
            "cat": "TỪ VỰNG: TOPPING & THANH TOÁN QR (2/3)",
            "word": "• cheese foam /ˈtʃiːz foʊm/ (n phr): lớp kem bọt phô mai béo ngậy mặn nhẹ",
            "ex": "  Ví dụ: Can I add cheese foam? (Cho em thêm ít kem cheese được không ạ?)"
        },
        {
            "cat": "TỪ VỰNG: TOPPING & THANH TOÁN QR (2/3)",
            "word": "• bank transfer /ˈbæŋk trænsfɜːr/ (n phr): chuyển khoản ngân hàng",
            "ex": "  Ví dụ: Do you take bank transfer? (Quán có nhận chuyển khoản không ạ?)"
        },
        {
            "cat": "TỪ VỰNG: TIỆN ÍCH & DỊCH VỤ PHỤC VỤ (3/3)",
            "word": "• straw /strɔː/ (n): ống hút",
            "ex": "  Ví dụ: Can I have a straw, please? (Cho em xin cái ống hút được không ạ?)"
        },
        {
            "cat": "TỪ VỰNG: TIỆN ÍCH & DỊCH VỤ PHỤC VỤ (3/3)",
            "word": "• upstairs /ˌʌpˈsterz/ (adv, adj, n): ở trên tầng / tầng trên",
            "ex": "  Ví dụ: Are there seats upstairs? (Có chỗ ngồi ở trên tầng không ạ?)"
        },
        {
            "cat": "TỪ VỰNG: TIỆN ÍCH & DỊCH VỤ PHỤC VỤ (3/3)",
            "word": "• air conditioner /ˈer kəndɪʃənər/ (AC /ˌeɪ ˈsiː/) (n): máy điều hòa nhiệt độ",
            "ex": "  Ví dụ: Excuse me! The AC is a bit cold. (Điều hòa hơi lạnh chị ạ.)"
        },
        {
            "cat": "TỪ VỰNG: TIỆN ÍCH & DỊCH VỤ PHỤC VỤ (3/3)",
            "word": "• napkin /ˈnæpkɪn/ (n): khăn giấy ăn | wipe /waɪp/ (v): lau chùi bàn",
            "ex": "  Ví dụ: Let me wipe the table and get you some napkins. (Để em lau bàn và lấy khăn giấy cho nha.)"
        }
    ]

    for v in vocab_items:
        add_vocab_slide(prs, v["cat"], v["word"], v["ex"])

    # Conversation 1
    c1_en = [
        "Barista: Hello! What can I get for you today?",
        "Customer: What's your best-selling drink?",
        "Barista: Definitely our iced coffee with milk!",
        "Customer: Awesome, I'll take that then!",
        "Barista: Medium or large?",
        "Customer: Large, please. Less ice and no sugar, please.",
        "Barista: For here or to go?",
        "Customer: To go, please.",
        "Barista: Would you like a pastry with that?",
        "Customer: No, thanks. That's all.",
        "Barista: Let me repeat your order: one large iced coffee with milk, less ice, no sugar, to go, right?",
        "Customer: Yes, that's right!"
    ]
    c1_vi = [
        "Nhân viên: Em chào chị ạ! Mình dùng gì ạ?",
        "Khách hàng: Món nào hot nhất đấy chị?",
        "Nhân viên: Cà phê sữa đá của quán em là hot nhất ạ!",
        "Khách hàng: Hay quá, cho em một ly đó nha!",
        "Nhân viên: Mình muốn size vừa hay size lớn ạ?",
        "Khách hàng: Cho em size lớn nhé. Ít đá và không bỏ đường nha.",
        "Nhân viên: Mình muốn dùng tại quán hay mang đi ạ?",
        "Khách hàng: Cho em mang đi.",
        "Nhân viên: Mình có muốn thêm bánh ngọt không ạ?",
        "Khách hàng: Dạ thôi, em cảm ơn ạ.",
        "Nhân viên: Để em nhắc lại order của mình nhé: một cà phê sữa đá size lớn, ít đá, không đường, mang đi, đúng không ạ?",
        "Khách hàng: Dạ đúng rồi ạ!"
    ]
    add_content_slide(prs, "CONVERSATION 1: ORDERING AT THE COUNTER", "Tình huống: Gọi món tại quầy, chọn size lớn, ít đá, không đường, mang đi.", c1_en, font_size=13)
    add_content_slide(prs, "HỘI THOẠI 1: BẢN DỊCH TIẾNG VIỆT", "Tình huống: Gọi món tại quầy & Tùy chỉnh khẩu vị", c1_vi, font_size=13)

    # Conversation 2
    c2_en = [
        "Barista: Have you decided yet?",
        "Customer: Can I get the one from the menu picture? Is this drink good?",
        "Barista: It's really good! Very refreshing.",
        "Customer: What toppings do you have? Can I add boba?",
        "Barista: We have boba and cheese foam!",
        "Customer: Can I add cheese foam too? And I'll have the same as her for my friend.",
        "Barista: Got it, two drinks with toppings! Cash or card?",
        "Customer: Do you take bank transfer? I'll pay now.",
        "Barista: Yes, we do! You can scan the QR code right here on the counter.",
        "Customer: Done! Should I wait here or take a seat?",
        "Barista: Here's your receipt. We're a bit busy right now, so there might be a short wait. Your drink will be ready in about five minutes.",
        "Customer: Please, let me know when it's ready."
    ]
    c2_vi = [
        "Nhân viên: Mình đã chọn được món chưa ạ?",
        "Khách hàng: Cho em một ly như trên hình menu này đi. Món này có ngon không chị?",
        "Nhân viên: Ngon lắm chị ơi, thanh mát dễ uống lắm!",
        "Khách hàng: Có topping gì đấy ạ? Cho em thêm trân châu nhá.",
        "Nhân viên: Quán có trân châu và kem cheese ạ!",
        "Khách hàng: Cho em thêm ít kem cheese được không ạ? Và cho em giống chị ấy cho bạn em nữa nha.",
        "Nhân viên: Dạ được ạ, hai ly thêm topping! Mình thanh toán tiền mặt hay chuyển khoản/thẻ ạ?",
        "Khách hàng: Có nhận chuyển khoản không ạ? Cho em thanh toán luôn nhá.",
        "Nhân viên: Dạ có ạ! Chị quét mã QR ngay trên quầy này nhé.",
        "Khách hàng: Em quét xong rồi ạ! Em đợi ở đây hay ra bàn ngồi ạ?",
        "Nhân viên: Hóa đơn của mình đây ạ. Quán hôm nay hơi đông nên chị đợi em một xíu nhé, khoảng 5 phút là có đồ ạ.",
        "Khách hàng: Khi nào xong thì gọi em nha!"
    ]
    add_content_slide(prs, "CONVERSATION 2: TOPPINGS & QR PAYMENT", "Tình huống: Thêm topping trân châu, kem cheese, chuyển khoản quét mã QR.", c2_en, font_size=12)
    add_content_slide(prs, "HỘI THOẠI 2: BẢN DỊCH TIẾNG VIỆT", "Tình huống: Topping, Thanh toán chuyển khoản QR & Chờ nhận đồ", c2_vi, font_size=12)

    # Conversation 3
    c3_en = [
        "Barista: Here's your drink. Enjoy your drink!",
        "Customer: Thank you! Can I have a straw, please? And extra ice, please.",
        "Barista: Here's your straw and extra ice for you!",
        "Customer: Are there seats upstairs?",
        "Barista: Yes, there are plenty of tables upstairs near the window!",
        "Customer: Great! What's the Wi-Fi password?",
        "Barista: The Wifi password is on the receipt.",
        "Customer: Excuse me! The AC up there is a bit cold.",
        "Barista: Let me turn up the AC for you! And let me wipe the table for you too. I'll get you some napkins.",
        "Customer: Thank you so much!",
        "Barista: Do you need anything else?",
        "Customer: No, that's perfect. Thank you!",
        "Barista: Thank you! See you again."
    ]
    c3_vi = [
        "Nhân viên: Em gửi chị đồ uống ạ. Chị dùng ngon miệng nha!",
        "Khách hàng: Cảm ơn em! Cho em xin cái ống hút và cho thêm tí đá nha.",
        "Nhân viên: Ống hút và đá thêm của mình đây ạ!",
        "Khách hàng: Có chỗ ngồi ở trên tầng không ạ?",
        "Nhân viên: Dạ có, trên tầng còn nhiều bàn gần cửa sổ lắm ạ!",
        "Khách hàng: Tuyệt quá! Pass Wi-Fi ở đây là gì đấy ạ?",
        "Nhân viên: Pass Wi-Fi ở trên hóa đơn ấy ạ.",
        "Khách hàng: Điều hòa trên đó hơi lạnh em ạ.",
        "Nhân viên: Để em chỉnh điều hòa ấm lên một chút ạ! Và để em lau bàn cho mình nha, em lấy thêm khăn giấy cho ạ.",
        "Khách hàng: Cảm ơn em nhiều nha!",
        "Nhân viên: Chị có cần gì nữa không ạ?",
        "Khách hàng: Dạ thế là tuyệt rồi, cảm ơn em!",
        "Nhân viên: Em cảm ơn nha! Lần sau lại ghé nha."
    ]
    add_content_slide(prs, "CONVERSATION 3: PICK-UP & AMENITIES", "Tình huống: Nhận đồ, xin ống hút, thêm đá, tìm chỗ tầng trên, pass Wi-Fi, máy lạnh.", c3_en, font_size=13)
    add_content_slide(prs, "HỘI THOẠI 3: BẢN DỊCH TIẾNG VIỆT", "Tình huống: Nhận đồ, Chỗ ngồi, Wi-Fi & Tiện ích quán", c3_vi, font_size=13)

    # Staff Phrases (Part 1 & 2)
    staff_p1 = [
        "1. Hello! What can I get for you today? (Em chào chị ạ! Mình dùng gì ạ?)",
        "2. What would you like to order? (Chị muốn gọi món gì đây ạ?)",
        "3. Have you decided yet? (Mình đã chọn được món chưa?)",
        "4. Medium or large? (Mình muốn size vừa hay là size lớn ạ?)",
        "5. How much sugar would you like? (Bao nhiêu đường ạ?)",
        "6. For here or to go? (Mình muốn dùng tại quán hay là mang đi ạ?)",
        "7. Would you like a pastry with that? (Mình có muốn thêm bánh không ạ?)",
        "8. Let me repeat your order. (Để em nhắc lại order của mình nhé.)",
        "9. One large iced coffee with milk, right? (Một cà phê sữa đá size lớn, đúng không ạ?)",
        "10. Cash or card? (Mình thanh toán tiền mặt hay chuyển khoản ạ?)",
        "11. Here's your receipt. (Hóa đơn của mình đây ạ.)"
    ]
    add_content_slide(prs, "MẪU CÂU NHÂN VIÊN QUÁN (STAFF - 1/2)", "Phần 1: Chào hỏi, Gợi ý & Xử lý đơn hàng", staff_p1, font_size=13)

    staff_p2 = [
        "12. We're a bit busy right now, so there might be a short wait. (Quán hôm nay hơi đông nên là chị đợi em 1 xíu nhé.)",
        "13. Your drink will be ready in about five minutes. (Mình đợi đồ uống khoảng 5 phút nhé.)",
        "14. Here's your drink. (Em gửi chị đồ uống ạ.)",
        "15. Careful! It's hot. (Cẩn thận nóng nha.)",
        "16. Enjoy your drink! (Chị dùng ngon miệng nha.)",
        "17. I'll get you some napkins. (Để em lấy khăn giấy cho ạ.)",
        "18. Let me wipe the table for you. (Để em lau bàn cho mình nha.)",
        "19. The Wifi password is on the receipt. (Pass Wi-Fi ở trên hóa đơn ấy ạ.)",
        "20. Do you need anything else? (Chị có cần gì nữa không ạ?)",
        "21. Thank you! See you again. (Em cảm ơn nha! Lần sau lại ghé nha.)"
    ]
    add_content_slide(prs, "MẪU CÂU NHÂN VIÊN QUÁN (STAFF - 2/2)", "Phần 2: Phục vụ, Tiện ích & Chăm sóc khách", staff_p2, font_size=13)

    # Customer Phrases (Part 1 & 2)
    cust_p1 = [
        "1. Can I get an iced coffee with milk? (Cho em một ly cà phê sữa đá.)",
        "2. Less ice, please. (Mà ít đá thôi ạ.)",
        "3. Extra ice, please. (Cho thêm tí đá nha.)",
        "4. No sugar, please. (Không bỏ đường ạ.)",
        "5. To go, please. (Cho em mang đi.)",
        "6. For here, please. (Em uống tại đây luôn ạ.)",
        "7. Is this drink good? (Chị ơi, món này có ngon không?)",
        "8. I'll have the same as her. (Cho em giống chị ấy nhá.)",
        "9. What toppings do you have? (Có topping gì đấy ạ?)",
        "10. Can I add boba? (Cho em thêm trân châu nhá.)",
        "11. Can I get the one from the menu picture? (Cho em một ly như trên hình menu này đi.)"
    ]
    add_content_slide(prs, "MẪU CÂU KHÁCH HÀNG (CUSTOMER - 1/2)", "Phần 1: Gọi món, Khẩu vị & Topping", cust_p1, font_size=13)

    cust_p2 = [
        "12. What's your best-selling drink? (Món nào hot nhất đấy chị?)",
        "13. Can I add cheese foam? (Cho em thêm ít kem cheese được không ạ?)",
        "14. I'll pay now. (Cho em thanh toán luôn nhá.)",
        "15. Do you take bank transfer? (Có nhận chuyển khoản không ạ?)",
        "16. Where do I scan the QR code? (Quét mã QR ở đâu ạ?)",
        "17. Should I wait here or take a seat? (Em đợi ở đây hay là ra bàn ngồi ạ?)",
        "18. Please, let me know when it's ready. (Khi nào xong thì gọi em nha!)",
        "19. Can I have a straw, please? (Cho em xin cái ống hút được không ạ?)",
        "20. Are there seats upstairs? (Có chỗ ngồi ở trên tầng không ạ?)",
        "21. What's the Wi-Fi password? (Pass Wi-Fi ở đây là gì đấy ạ?)",
        "22. Excuse me! The AC is a bit cold. (Điều hòa hơi lạnh chị ạ.)"
    ]
    add_content_slide(prs, "MẪU CÂU KHÁCH HÀNG (CUSTOMER - 2/2)", "Phần 2: Thanh toán QR, Chỗ ngồi & Tiện ích", cust_p2, font_size=13)

    # Slide Summary
    summary_lines = [
        "1. CÔNG THỨC GỌI MÓN TIÊU CHUẨN:",
        "   Can I get a [Size] [Temperature: Hot / Iced] [Drink Name] with [Type of Milk]?",
        "   -> Can I get a large iced latte with oat milk, please?",
        "",
        "2. CÔNG THỨC TÙY CHỈNH (CUSTOMIZING):",
        "   - With [less ice / extra ice / no ice]",
        "   - With [less sugar / no sugar / extra shot]",
        "   - Add [boba / cheese foam / jelly]",
        "",
        "3. CÔNG THỨC HỎI TIỆN ÍCH & PHỤC VỤ:",
        "   - Do you take [bank transfer / card / cash]?",
        "   - Can I have [a straw / extra napkins], please?",
        "   - Is there [Wi-Fi / a seat upstairs / a power outlet]?"
    ]
    add_content_slide(prs, "SUMMARY & ORDERING CHEAT SHEET", "Tổng kết công thức gọi món phản xạ nhanh", summary_lines, font_size=14)

    try:
        prs.save(OUTPUT_PATH)
        print(f"SUCCESS: Saved {len(prs.slides)} slides to {OUTPUT_PATH}")
    except PermissionError:
        fallback_path = os.path.join(PROJECT_DIR, "materials", "slides", "Cafe_v2.pptx")
        prs.save(fallback_path)
        print(f"NOTE: {OUTPUT_PATH} is currently open in PowerPoint. Saved {len(prs.slides)} slides to {fallback_path}")

if __name__ == "__main__":
    main()
