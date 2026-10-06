import sys
import shutil
from pptx import Presentation
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls

sys.stdout.reconfigure(encoding='utf-8')

# Backup original file
shutil.copyfile("Airport.pptx", "Airport_backup.pptx")
print("Backed up Airport.pptx -> Airport_backup.pptx")

prs = Presentation("Airport.pptx")
print(f"Original slide count: {len(prs.slides)}")

vocab_slides = [
    # --- AIRPORT (3/3) ---
    {
        "category": "TỪ VỰNG: THỦ TỤC VÀ DỊCH VỤ KHÁC TẠI SÂN BAY (3/3)",
        "word": "• baggage claim /ˈbæɡɪdʒ kleɪm/ (n): khu vực nhận hành lý",
        "example": "  Ví dụ: The baggage claim area is on the first floor."
    },
    {
        "category": "TỪ VỰNG: THỦ TỤC VÀ DỊCH VỤ KHÁC TẠI SÂN BAY (3/3)",
        "word": "• conveyor belt /kənˈveɪər belt/ (n): băng chuyền",
        "example": "  Ví dụ: Wait for your luggage at the conveyor belt."
    },
    {
        "category": "TỪ VỰNG: THỦ TỤC VÀ DỊCH VỤ KHÁC TẠI SÂN BAY (3/3)",
        "word": "• customs /ˈkʌstəmz/ (n): hải quan",
        "example": "  Ví dụ: You must go through customs when entering the country."
    },
    {
        "category": "TỪ VỰNG: THỦ TỤC VÀ DỊCH VỤ KHÁC TẠI SÂN BAY (3/3)",
        "word": "• immigration /ˌɪmɪˈɡreɪʃən/ (n): cửa kiểm soát xuất nhập cảnh",
        "example": "  Ví dụ: The immigration line was very long today."
    },
    {
        "category": "TỪ VỰNG: THỦ TỤC VÀ DỊCH VỤ KHÁC TẠI SÂN BAY (3/3)",
        "word": "• declaration form /ˌdekləˈreɪʃən fɔːrm/ (n): tờ khai hải quan",
        "example": "  Ví dụ: Please fill out the declaration form before landing."
    },
    {
        "category": "TỪ VỰNG: THỦ TỤC VÀ DỊCH VỤ KHÁC TẠI SÂN BAY (3/3)",
        "word": "• duty-free shop /ˌdjuːti ˈfriː ʃɑːp/ (n): cửa hàng miễn thuế",
        "example": "  Ví dụ: I bought perfume at the duty-free shop."
    },
    {
        "category": "TỪ VỰNG: THỦ TỤC VÀ DỊCH VỤ KHÁC TẠI SÂN BAY (3/3)",
        "word": "• exchange booth /ɪksˈtʃeɪndʒ buːθ/ (n): quầy đổi tiền",
        "example": "  Ví dụ: There is an exchange booth near the exit."
    },

    # --- FLIGHT (1/2) ---
    {
        "category": "TỪ VỰNG: LỊCH TRÌNH VÀ LOẠI CHUYẾN BAY (1/2)",
        "word": "• flight /ˈflaɪt/ (n): chuyến bay",
        "example": "  Ví dụ: My flight departs at 3 PM."
    },
    {
        "category": "TỪ VỰNG: LỊCH TRÌNH VÀ LOẠI CHUYẾN BAY (1/2)",
        "word": "• departure /dɪˈpɑːrtʃər/ (n): sự khởi hành",
        "example": "  Ví dụ: The departure time has been changed to 4 PM."
    },
    {
        "category": "TỪ VỰNG: LỊCH TRÌNH VÀ LOẠI CHUYẾN BAY (1/2)",
        "word": "• arrival /əˈraɪvəl/ (n): sự đến nơi",
        "example": "  Ví dụ: Our estimated arrival time is 5:30 PM."
    },
    {
        "category": "TỪ VỰNG: LỊCH TRÌNH VÀ LOẠI CHUYẾN BAY (1/2)",
        "word": "• destination /ˌdestɪˈneɪʃən/ (n): điểm đến",
        "example": "  Ví dụ: What is your final destination?"
    },
    {
        "category": "TỪ VỰNG: LỊCH TRÌNH VÀ LOẠI CHUYẾN BAY (1/2)",
        "word": "• layover /ˈleɪˌoʊvər/ (n): quá cảnh",
        "example": "  Ví dụ: We have a 3-hour layover in Singapore."
    },
    {
        "category": "TỪ VỰNG: LỊCH TRÌNH VÀ LOẠI CHUYẾN BAY (1/2)",
        "word": "• connecting flight /kəˈnektɪŋ flaɪt/ (n): chuyến bay nối chuyến",
        "example": "  Ví dụ: I need to catch a connecting flight to Tokyo."
    },
    {
        "category": "TỪ VỰNG: LỊCH TRÌNH VÀ LOẠI CHUYẾN BAY (1/2)",
        "word": "• direct flight /dəˈrekt flaɪt/ (n): chuyến bay thẳng",
        "example": "  Ví dụ: Is this a direct flight to London?"
    },

    # --- FLIGHT (2/2) ---
    {
        "category": "TỪ VỰNG: HOẠT ĐỘNG VÀ TRẠNG THÁI CHUYẾN BAY (2/2)",
        "word": "• takeoff /ˈteɪkɔːf/ (n): sự cất cánh",
        "example": "  Ví dụ: Please fasten your seatbelt before takeoff."
    },
    {
        "category": "TỪ VỰNG: HOẠT ĐỘNG VÀ TRẠNG THÁI CHUYẾN BAY (2/2)",
        "word": "• landing /ˈlændɪŋ/ (n): sự hạ cánh",
        "example": "  Ví dụ: We had a very smooth landing."
    },
    {
        "category": "TỪ VỰNG: HOẠT ĐỘNG VÀ TRẠNG THÁI CHUYẾN BAY (2/2)",
        "word": "• board /bɔːrd/ (v): lên (máy bay)",
        "example": "  Ví dụ: Passengers in rows 20–30, please board now."
    },
    {
        "category": "TỪ VỰNG: HOẠT ĐỘNG VÀ TRẠNG THÁI CHUYẾN BAY (2/2)",
        "word": "• boarding /ˈbɔːrdɪŋ/ (n): việc lên máy bay",
        "example": "  Ví dụ: Boarding will begin at 2:30 PM."
    },
    {
        "category": "TỪ VỰNG: HOẠT ĐỘNG VÀ TRẠNG THÁI CHUYẾN BAY (2/2)",
        "word": "• delayed /dɪˈleɪd/ (adj): bị hoãn/trễ",
        "example": "  Ví dụ: The flight has been delayed by one hour."
    },
    {
        "category": "TỪ VỰNG: HOẠT ĐỘNG VÀ TRẠNG THÁI CHUYẾN BAY (2/2)",
        "word": "• cancelled /ˈkænsəld/ (adj): bị hủy",
        "example": "  Ví dụ: Unfortunately, your flight has been cancelled."
    },
    {
        "category": "TỪ VỰNG: HOẠT ĐỘNG VÀ TRẠNG THÁI CHUYẾN BAY (2/2)",
        "word": "• on time /ɑːn taɪm/ (adv): đúng giờ",
        "example": "  Ví dụ: Don't worry, the flight is on time."
    },
    {
        "category": "TỪ VỰNG: HOẠT ĐỘNG VÀ TRẠNG THÁI CHUYẾN BAY (2/2)",
        "word": "• turbulence /ˈtɜːrbjələns/ (n): nhiễu động không khí",
        "example": "  Ví dụ: We are experiencing some turbulence."
    },
    {
        "category": "TỪ VỰNG: HOẠT ĐỘNG VÀ TRẠNG THÁI CHUYẾN BAY (2/2)",
        "word": "• announcement /əˈnaʊnsmənt/ (n): thông báo",
        "example": "  Ví dụ: Please pay attention to the announcement."
    },
    {
        "category": "TỪ VỰNG: HOẠT ĐỘNG VÀ TRẠNG THÁI CHUYẾN BAY (2/2)",
        "word": "• runway /ˈrʌnweɪ/ (n): đường băng",
        "example": "  Ví dụ: The plane is waiting on the runway."
    },

    # --- ON THE PLANE (1/2) ---
    {
        "category": "TỪ VỰNG: CHỖ NGỒI VÀ TIỆN NGHI TRÊN MÁY BAY (1/2)",
        "word": "• aisle seat /ˈaɪl siːt/ (n): ghế ngồi phía lối đi",
        "example": "  Ví dụ: I prefer an aisle seat so I can stretch my legs."
    },
    {
        "category": "TỪ VỰNG: CHỖ NGỒI VÀ TIỆN NGHI TRÊN MÁY BAY (1/2)",
        "word": "• window seat /ˈwɪndoʊ siːt/ (n): ghế cạnh cửa sổ",
        "example": "  Ví dụ: Can I have a window seat, please?"
    },
    {
        "category": "TỪ VỰNG: CHỖ NGỒI VÀ TIỆN NGHI TRÊN MÁY BAY (1/2)",
        "word": "• middle seat /ˈmɪdl siːt/ (n): ghế giữa",
        "example": "  Ví dụ: Nobody likes the middle seat."
    },
    {
        "category": "TỪ VỰNG: CHỖ NGỒI VÀ TIỆN NGHI TRÊN MÁY BAY (1/2)",
        "word": "• overhead bin /ˌoʊvərˈhed bɪn/ (n): ngăn chứa hành lý phía trên",
        "example": "  Ví dụ: Please put your bag in the overhead bin."
    },
    {
        "category": "TỪ VỰNG: CHỖ NGỒI VÀ TIỆN NGHI TRÊN MÁY BAY (1/2)",
        "word": "• tray table /treɪ ˈteɪbəl/ (n): bàn ăn gập",
        "example": "  Ví dụ: Please fold up your tray table for landing."
    },
    {
        "category": "TỪ VỰNG: CHỖ NGỒI VÀ TIỆN NGHI TRÊN MÁY BAY (1/2)",
        "word": "• seat pocket /ˈsiːt ˈpɑːkɪt/ (n): túi ghế (phía trước)",
        "example": "  Ví dụ: The safety card is in the seat pocket."
    },
    {
        "category": "TỪ VỰNG: CHỖ NGỒI VÀ TIỆN NGHI TRÊN MÁY BAY (1/2)",
        "word": "• blanket /ˈblæŋkɪt/ (n): chăn, mền",
        "example": "  Ví dụ: Could I have a blanket, please? It's cold."
    },
    {
        "category": "TỪ VỰNG: CHỖ NGỒI VÀ TIỆN NGHI TRÊN MÁY BAY (1/2)",
        "word": "• pillow /ˈpɪloʊ/ (n): gối",
        "example": "  Ví dụ: Can I get an extra pillow?"
    },
    {
        "category": "TỪ VỰNG: CHỖ NGỒI VÀ TIỆN NGHI TRÊN MÁY BAY (1/2)",
        "word": "• headphones /ˈhedfoʊnz/ (n): tai nghe",
        "example": "  Ví dụ: The headphones are in the seat pocket in front of you."
    },

    # --- ON THE PLANE (2/2) ---
    {
        "category": "TỪ VỰNG: AN TOÀN VÀ PHỤC VỤ TRÊN MÁY BAY (2/2)",
        "word": "• seatbelt /ˈsiːtbelt/ (n): dây an toàn",
        "example": "  Ví dụ: Please fasten your seatbelt."
    },
    {
        "category": "TỪ VỰNG: AN TOÀN VÀ PHỤC VỤ TRÊN MÁY BAY (2/2)",
        "word": "• fasten /ˈfæsən/ (v): thắt, cài",
        "example": "  Ví dụ: Please fasten your seatbelt during takeoff."
    },
    {
        "category": "TỪ VỰNG: AN TOÀN VÀ PHỤC VỤ TRÊN MÁY BAY (2/2)",
        "word": "• emergency exit /ɪˈmɜːrdʒənsi ˈeɡzɪt/ (n): lối thoát hiểm",
        "example": "  Ví dụ: The emergency exit is located over the wing."
    },
    {
        "category": "TỪ VỰNG: AN TOÀN VÀ PHỤC VỤ TRÊN MÁY BAY (2/2)",
        "word": "• life vest /ˈlaɪf vest/ (n): áo phao cứu sinh",
        "example": "  Ví dụ: The life vest is under your seat."
    },
    {
        "category": "TỪ VỰNG: AN TOÀN VÀ PHỤC VỤ TRÊN MÁY BAY (2/2)",
        "word": "• beverage /ˈbevərɪdʒ/ (n): đồ uống",
        "example": "  Ví dụ: Would you like a beverage?"
    },
    {
        "category": "TỪ VỰNG: AN TOÀN VÀ PHỤC VỤ TRÊN MÁY BAY (2/2)",
        "word": "• meal /ˈmiːl/ (n): bữa ăn",
        "example": "  Ví dụ: The meal will be served in 30 minutes."
    },

    # --- TRANSPORTATION ---
    {
        "category": "TỪ VỰNG: PHƯƠNG TIỆN DI CHUYỂN (TRANSPORTATION)",
        "word": "• taxi /ˈtæksi/ (n): xe taxi",
        "example": "  Ví dụ: Let's take a taxi to the hotel."
    },
    {
        "category": "TỪ VỰNG: PHƯƠNG TIỆN DI CHUYỂN (TRANSPORTATION)",
        "word": "• subway /ˈsʌbweɪ/ (n): tàu điện ngầm",
        "example": "  Ví dụ: The subway is cheaper than a taxi."
    },
    {
        "category": "TỪ VỰNG: PHƯƠNG TIỆN DI CHUYỂN (TRANSPORTATION)",
        "word": "• shuttle bus /ˈʃʌtl bʌs/ (n): xe buýt đưa đón",
        "example": "  Ví dụ: There is a free shuttle bus to the city center."
    },
    {
        "category": "TỪ VỰNG: PHƯƠNG TIỆN DI CHUYỂN (TRANSPORTATION)",
        "word": "• rental car /ˈrentl kɑːr/ (n): xe thuê",
        "example": "  Ví dụ: I reserved a rental car at the airport."
    }
]

def add_p(tf, text, is_bold=False):
    if len(tf.paragraphs) == 1 and tf.paragraphs[0].text == '':
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.text = text
    pPr = p._p.get_or_add_pPr()
    pPr.set('marL', '0')
    pPr.set('indent', '0')
    ns = nsdecls('a')
    buNone = parse_xml(f'<a:buNone {ns}/>')
    pPr.append(buNone)
    if is_bold:
        for r in p.runs:
            r.font.bold = True
    return p

# Store original conversation slide elements (Slide 25, 26, 27)
sldIdLst = prs.element.sldIdLst
conv_elements = [sldIdLst[24], sldIdLst[25], sldIdLst[26]]

# Add each vocabulary slide
for item in vocab_slides:
    slide = prs.slides.add_slide(prs.slide_layouts[1]) # Title and Content
    title_shape = slide.shapes.title
    title_shape.text = "Vocabulary"
    for r in title_shape.text_frame.paragraphs[0].runs:
        r.font.bold = True
    
    body = slide.shapes.placeholders[1]
    body.left = 838200
    body.top = 1265125
    body.width = 6614160
    body.height = 4911838
    
    tf = body.text_frame
    tf.word_wrap = True
    
    add_p(tf, item["category"], is_bold=True)
    add_p(tf, "")
    add_p(tf, item["word"])
    add_p(tf, item["example"])

# Move conversation slides to the end
for el in conv_elements:
    sldIdLst.append(el)

prs.save("Airport.pptx")
print(f"Successfully populated slides! Total slides now: {len(prs.slides)}")
