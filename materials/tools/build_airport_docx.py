# -*- coding: utf-8 -*-
"""
Script to generate a comprehensive, professional Microsoft Word document (airport.docx)
from airport.md, adding Synonyms and Collocations for every vocabulary word,
and presenting both Enriched Vocabulary and 12 Practical Bilingual Dialogues.
"""

import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

# --- VOCABULARY ENRICHMENT DATABASE (66 WORDS) ---
ENRICHED_DATA = {
    # 👤 People / Con người (8 words)
    "passenger": {
        "synonyms": [
            ("traveler", "người đi lại / du khách"),
            ("commuter", "người đi lại hàng ngày / người đi vé tháng"),
            ("flyer", "khách đi máy bay (vd: frequent flyer)"),
            ("voyager", "lữ khách, người thực hiện chuyến đi dài (văn phong trang trọng)")
        ],
        "collocations": [
            ("transit passenger", "hành khách quá cảnh"),
            ("fellow passenger", "hành khách đi cùng chuyến"),
            ("passenger manifest", "danh sách hành khách trên chuyến bay"),
            ("passenger terminal", "nhà ga hành khách"),
            ("frequent flyer", "hành khách bay thường xuyên (hội viên tích dặm)")
        ],
        "example_vn": "Hành khách đã xuất trình thẻ lên máy bay tại cửa khởi hành."
    },
    "flight attendant": {
        "synonyms": [
            ("cabin crew member", "thành viên đoàn tiếp viên"),
            ("air steward / air stewardess", "nam / nữ tiếp viên (cách gọi truyền thống)"),
            ("in-flight crew", "đội ngũ phục vụ trên khoang")
        ],
        "collocations": [
            ("call the flight attendant", "bấm chuông gọi tiếp viên"),
            ("senior flight attendant / purser", "tiếp viên trưởng"),
            ("flight attendant call button", "nút gọi tiếp viên trên trần máy bay"),
            ("listen to the flight attendant", "lắng nghe hướng dẫn của tiếp viên")
        ],
        "example_vn": "Tiếp viên hàng không đã phục vụ đồ uống trong suốt chuyến bay."
    },
    "stewardess": {
        "synonyms": [
            ("female flight attendant", "nữ tiếp viên hàng không (thuật ngữ hiện đại)"),
            ("air hostess", "nữ tiếp viên (từ truyền thống, hiện ít dùng hơn)")
        ],
        "collocations": [
            ("chief stewardess", "nữ tiếp viên trưởng"),
            ("airline stewardess", "nữ tiếp viên của hãng hàng không"),
            ("ground hostess", "nữ nhân viên mặt đất")
        ],
        "example_vn": "Cô tiếp viên đã hướng dẫn tôi tìm đúng chỗ ngồi."
    },
    "steward": {
        "synonyms": [
            ("male flight attendant", "nam tiếp viên hàng không (thuật ngữ trung tính)"),
            ("air steward", "nam tiếp viên hàng không"),
            ("cabin attendant", "nhân viên phục vụ khoang khách")
        ],
        "collocations": [
            ("chief steward", "nam tiếp viên trưởng"),
            ("cabin steward", "tiếp viên phụ trách khoang"),
            ("ground steward", "nam nhân viên phục vụ mặt đất")
        ],
        "example_vn": "Nam tiếp viên đã mang thêm chăn đắp cho chúng tôi."
    },
    "cabin crew": {
        "synonyms": [
            ("flight attendants", "đội ngũ tiếp viên"),
            ("in-flight team", "đội ngũ phục vụ trên chuyến bay"),
            ("air crew", "phi hành đoàn (thường bao gồm cả phi công)")
        ],
        "collocations": [
            ("member of the cabin crew", "thành viên đoàn tiếp viên"),
            ("cabin crew instructions", "chỉ dẫn an toàn của đoàn tiếp viên"),
            ("cabin crew safety demonstration", "màn minh họa an toàn bay"),
            ("alert the cabin crew", "báo động cho đoàn tiếp viên khi có sự cố")
        ],
        "example_vn": "Đoàn tiếp viên đã chuẩn bị khoang hành khách sẵn sàng để cất cánh."
    },
    "pilot": {
        "synonyms": [
            ("aviator", "nhà hàng không / phi công"),
            ("flyer", "người lái máy bay"),
            ("airman", "phi công / quân nhân không quân")
        ],
        "collocations": [
            ("commercial pilot", "phi công hàng không dân dụng"),
            ("airline pilot", "phi công hãng hàng không"),
            ("pilot in command (PIC)", "phi công chỉ huy chuyến bay"),
            ("automatic pilot / autopilot", "chế độ lái tự động"),
            ("pilot's announcement", "thông báo của phi công qua loa")
        ],
        "example_vn": "Phi công thông báo chúng tôi sẽ hạ cánh trong 20 phút nữa."
    },
    "captain": {
        "synonyms": [
            ("aircraft commander", "chỉ huy trưởng tàu bay"),
            ("pilot-in-command (PIC)", "cơ trưởng / phi công chỉ huy"),
            ("skipper", "thuyền trưởng / cơ trưởng (tiếng lóng thân mật trong ngành)")
        ],
        "collocations": [
            ("airline captain", "cơ trưởng hãng hàng không"),
            ("this is your captain speaking", "đây là cơ trưởng đang nói (lời mở đầu phát thanh)"),
            ("under the captain's command", "dưới quyền điều hành của cơ trưởng"),
            ("senior captain", "cơ trưởng thâm niên / kỳ cựu")
        ],
        "example_vn": "Cơ trưởng đã gửi lời chào mừng toàn thể hành khách lên máy bay."
    },
    "co-pilot": {
        "synonyms": [
            ("first officer (FO)", "cơ phó (thuật ngữ chính thức quốc tế)"),
            ("second-in-command (SIC)", "người chỉ huy thứ hai trong buồng lái")
        ],
        "collocations": [
            ("act as co-pilot", "đảm nhận vị trí cơ phó"),
            ("first officer / co-pilot duties", "nhiệm vụ của cơ phó"),
            ("assist the captain", "hỗ trợ cơ trưởng vận hành máy bay"),
            ("co-pilot controls", "bộ cần điều khiển của cơ phó")
        ],
        "example_vn": "Cơ phó đã hỗ trợ cơ trưởng trong suốt quá trình hạ cánh."
    },

    # 🏢 Airport / Sân bay (22 words)
    "airport": {
        "synonyms": [
            ("airfield", "sân bay nhỏ / phi trường quân sự"),
            ("aerodrome", "phi trường (từ truyền thống Anh-Anh)"),
            ("airstrip", "đường băng dã chiến ngắn")
        ],
        "collocations": [
            ("international / domestic airport", "sân bay quốc tế / nội địa"),
            ("airport terminal", "nhà ga sân bay"),
            ("airport shuttle", "xe buýt trung chuyển sân bay"),
            ("airport code", "mã sân bay (IATA: SGN, HAN, BKK...)"),
            ("head to / arrive at the airport", "khởi hành đến / tới sân bay")
        ],
        "example_vn": "Chúng tôi đã đến sân bay sớm hai tiếng đồng hồ."
    },
    "terminal": {
        "synonyms": [
            ("concourse", "khu sảnh / hành lang ga sân bay"),
            ("passenger terminal", "nhà ga hành khách"),
            ("air terminal", "nhà ga hàng không")
        ],
        "collocations": [
            ("international terminal (T2)", "nhà ga quốc tế"),
            ("domestic terminal (T1)", "nhà ga nội địa"),
            ("terminal building", "tòa nhà ga"),
            ("transfer between terminals", "di chuyển giữa các nhà ga"),
            ("departure / arrival terminal", "nhà ga đi / nhà ga đến")
        ],
        "example_vn": "Chuyến bay của chúng tôi khởi hành từ Nhà ga số 2."
    },
    "check-in counter": {
        "synonyms": [
            ("check-in desk", "bàn làm thủ tục"),
            ("ticket counter", "quầy vé và thủ tục")
        ],
        "collocations": [
            ("go to / proceed to the check-in counter", "đi tới quầy làm thủ tục"),
            ("open / close the check-in counter", "mở / đóng quầy làm thủ tục"),
            ("queue at the check-in counter", "xếp hàng chờ tại quầy thủ tục"),
            ("self-service check-in kiosk", "quầy làm thủ tục tự động qua máy cảm ứng")
        ],
        "example_vn": "Vui lòng tới quầy làm thủ tục để nhận thẻ lên máy bay."
    },
    "gate": {
        "synonyms": [
            ("boarding gate", "cổng lên máy bay"),
            ("departure gate", "cửa khởi hành")
        ],
        "collocations": [
            ("gate number", "số cổng khởi hành"),
            ("gate closes at...", "cổng sẽ đóng lúc..."),
            ("gate change", "thay đổi cổng khởi hành"),
            ("proceed to gate", "di chuyển ra cửa lên tàu bay"),
            ("gate agent", "nhân viên phụ trách tại cửa khởi hành")
        ],
        "example_vn": "Chuyến bay của bạn khởi hành từ cổng 15A."
    },
    "security check": {
        "synonyms": [
            ("security screening", "quá trình soi chiếu an ninh"),
            ("security checkpoint", "trạm kiểm tra an ninh"),
            ("TSA checkpoint", "cửa kiểm tra an ninh TSA (đặc thù tại Mỹ)")
        ],
        "collocations": [
            ("go through security check", "đi qua cổng kiểm tra an ninh"),
            ("clear security check", "hoàn tất thủ tục kiểm tra an ninh"),
            ("strict security check", "kiểm tra an ninh nghiêm ngặt"),
            ("security check line / queue", "hàng người chờ kiểm tra an ninh")
        ],
        "example_vn": "Việc kiểm tra an ninh mất khoảng 15 phút."
    },
    "boarding pass": {
        "synonyms": [
            ("boarding ticket", "vé / thẻ lên máy bay"),
            ("flight coupon", "cuống vé chuyến bay"),
            ("e-boarding pass / mobile boarding pass", "thẻ lên máy bay điện tử trên điện thoại")
        ],
        "collocations": [
            ("issue a boarding pass", "in / cấp thẻ lên tàu bay"),
            ("show / present your boarding pass", "xuất trình thẻ lên máy bay"),
            ("scan the boarding pass", "quét mã thẻ lên máy bay tại cổng"),
            ("print out the boarding pass", "in thẻ lên máy bay giấy")
        ],
        "example_vn": "Đừng để quên thẻ lên máy bay tại quầy làm thủ tục."
    },
    "passport": {
        "synonyms": [
            ("travel document", "giấy tờ đi lại quốc tế"),
            ("identification document (ID)", "giấy tờ tùy thân")
        ],
        "collocations": [
            ("valid / expired passport", "hộ chiếu còn hạn / hết hạn"),
            ("biometric / e-passport", "hộ chiếu gắn chip điện tử"),
            ("passport control", "khu vực kiểm soát hộ chiếu"),
            ("renew a passport", "gia hạn / làm mới hộ chiếu"),
            ("hold a passport", "sở hữu hộ chiếu quốc gia nào")
        ],
        "example_vn": "Làm ơn cho tôi xem hộ chiếu của bạn được không?"
    },
    "visa": {
        "synonyms": [
            ("entry permit", "giấy phép nhập cảnh"),
            ("travel authorization", "ủy quyền đi lại (vd: ESTA, ETA)")
        ],
        "collocations": [
            ("tourist / business / transit visa", "thị thực du lịch / công tác / quá cảnh"),
            ("apply for a visa", "nộp đơn xin cấp thị thực"),
            ("grant / issue a visa", "cấp thị thực"),
            ("visa on arrival (VoA)", "thị thực cấp tại cửa khẩu"),
            ("visa exemption / visa-free", "miễn thị thực")
        ],
        "example_vn": "Bạn có thị thực còn hiệu lực đối với quốc gia này không?"
    },
    "luggage": {
        "synonyms": [
            ("baggage", "hành lý (từ đồng nghĩa phổ biến trong tiếng Mỹ)"),
            ("bags", "túi xách, túi hành lý"),
            ("suitcases", "vali"),
            ("belongings", "đồ đạc cá nhân")
        ],
        "collocations": [
            ("piece of luggage", "kiện hành lý (lưu ý: luggage không đếm được)"),
            ("excess luggage", "hành lý vượt cước"),
            ("lost / misplaced luggage", "hành lý thất lạc"),
            ("luggage tag", "thẻ gắn định danh hành lý"),
            ("luggage trolley / cart", "xe đẩy hành lý")
        ],
        "example_vn": "Bạn có tổng cộng bao nhiêu kiện hành lý?"
    },
    "baggage": {
        "synonyms": [
            ("luggage", "hành lý (phổ biến trong tiếng Anh-Anh)"),
            ("belongings", "tài sản cá nhân mang theo"),
            ("gear", "đồ đạc hành lý")
        ],
        "collocations": [
            ("baggage allowance", "định mức hành lý miễn cước được phép mang"),
            ("baggage drop-off", "quầy ký gửi hành lý nhanh"),
            ("baggage handler", "nhân viên bốc dỡ hành lý sân bay"),
            ("excess baggage fee", "phí phạt quá cước hành lý"),
            ("left baggage office", "quầy giữ hành lý tạm thời")
        ],
        "example_vn": "Hành lý của bạn sẽ được chuyển thẳng đến điểm đến cuối cùng."
    },
    "carry-on bag": {
        "synonyms": [
            ("hand luggage", "hành lý xách tay (Anh-Anh)"),
            ("cabin bag / cabin baggage", "hành lý xách tay mang lên khoang khách"),
            ("carry-on luggage", "hành lý xách tay")
        ],
        "collocations": [
            ("carry-on bag size / dimensions", "kích thước chuẩn của hành lý xách tay"),
            ("bring a carry-on bag", "mang một kiện hành lý xách tay"),
            ("stow a carry-on bag", "cất gọn hành lý xách tay lên hộc trên trần"),
            ("carry-on weight limit", "giới hạn cân nặng hành lý xách tay (thường 7kg)")
        ],
        "example_vn": "Bạn được phép mang theo một túi hành lý xách tay lên máy bay."
    },
    "checked luggage": {
        "synonyms": [
            ("checked baggage", "hành lý ký gửi"),
            ("hold luggage / hold baggage", "hành lý gửi trong khoang hàng (Anh-Anh)")
        ],
        "collocations": [
            ("check in luggage", "làm thủ tục ký gửi hành lý"),
            ("checked luggage weight limit", "giới hạn trọng lượng hành lý ký gửi (vd 23kg, 30kg)"),
            ("drop off checked luggage", "giao hành lý ký gửi tại quầy"),
            ("retrieve checked luggage", "nhận lại hành lý ký gửi sau chuyến bay")
        ],
        "example_vn": "Mỗi hành khách được phép ký gửi một kiện hành lý."
    },
    "scale": {
        "synonyms": [
            ("weighing machine", "máy cân trọng lượng"),
            ("weighing scales", "bàn cân"),
            ("balance", "cái cân (cân đối trọng)")
        ],
        "collocations": [
            ("luggage scale", "cân hành lý"),
            ("digital scale", "cân điện tử hiển thị số"),
            ("place on the scale", "đặt kiện hàng lên bàn cân"),
            ("read the scale", "đọc chỉ số trên cân")
        ],
        "example_vn": "Xin vui lòng đặt túi của bạn lên bàn cân."
    },
    "overweight": {
        "synonyms": [
            ("exceeding weight limit", "vượt quá trọng lượng quy định"),
            ("heavy", "quá nặng"),
            ("oversized", "quá khổ về kích thước")
        ],
        "collocations": [
            ("overweight luggage / baggage", "hành lý vượt quá cân quy định"),
            ("overweight fee / surcharge", "phí phụ thu hành lý quá cân"),
            ("overweight by [number] kilos", "quá cân bao nhiêu kilogam")
        ],
        "example_vn": "Hành lý của bạn bị quá cân 3 kilôgam."
    },
    "baggage claim": {
        "synonyms": [
            ("baggage reclaim", "khu nhận lại hành lý (chuẩn Anh-Anh)"),
            ("luggage reclaim", "nơi lấy lại hành lý"),
            ("baggage collection area", "khu vực tập kết trả hành lý")
        ],
        "collocations": [
            ("baggage claim area / hall", "khu vực / sảnh trả hành lý"),
            ("baggage claim tag / check", "phiếu / cuống vé đối chiếu nhận hành lý"),
            ("proceed to baggage claim", "di chuyển đến khu nhận hành lý"),
            ("baggage claim belt", "băng chuyền tại khu trả hành lý")
        ],
        "example_vn": "Khu vực nhận hành lý nằm ở tầng một."
    },
    "conveyor belt": {
        "synonyms": [
            ("baggage carousel", "băng chuyền tròn trả hành lý"),
            ("luggage belt", "băng tải hành lý"),
            ("carousel", "băng chuyền hành lý quay vòng")
        ],
        "collocations": [
            ("luggage conveyor belt", "băng tải chuyển hành lý"),
            ("wait at the conveyor belt", "chờ đợi cạnh băng chuyền"),
            ("come off the conveyor belt", "được đẩy ra từ băng chuyền"),
            ("grab a bag off the conveyor belt", "nhấc vali ra khỏi băng chuyền")
        ],
        "example_vn": "Hãy chờ hành lý của bạn tại băng chuyền."
    },
    "customs": {
        "synonyms": [
            ("customs and border protection (CBP)", "cục hải quan và biên phòng"),
            ("border inspection", "khu kiểm soát biên giới"),
            ("tariff bureau", "cơ quan thuế quan")
        ],
        "collocations": [
            ("go through / clear customs", "làm thủ tục thông quan / qua cửa hải quan"),
            ("customs officer", "cán bộ hải quan"),
            ("customs declaration", "khai báo hàng hóa hải quan"),
            ("declare goods at customs", "khai báo hàng thuộc diện chịu thuế"),
            ("customs duty / duty-free", "thuế hải quan / miễn thuế hải quan")
        ],
        "example_vn": "Bạn phải đi qua cửa hải quan khi nhập cảnh vào quốc gia này."
    },
    "immigration": {
        "synonyms": [
            ("passport control", "khu kiểm soát hộ chiếu xuất nhập cảnh"),
            ("border control", "cơ quan kiểm soát biên giới"),
            ("immigration checkpoint", "trạm kiểm soát nhập cảnh")
        ],
        "collocations": [
            ("pass through immigration", "làm thủ tục kiểm tra nhập cảnh"),
            ("immigration officer", "cán bộ quản lý xuất nhập cảnh"),
            ("immigration clearance", "thủ tục chấp thuận nhập cảnh"),
            ("immigration form", "mẫu đơn khai báo nhập cảnh")
        ],
        "example_vn": "Hàng người chờ kiểm tra xuất nhập cảnh hôm nay rất dài."
    },
    "declaration form": {
        "synonyms": [
            ("customs form", "tờ khai hải quan"),
            ("landing card", "thẻ đến / tờ khai nhập cảnh (Anh-Anh)"),
            ("arrival card", "phiếu nhập cảnh")
        ],
        "collocations": [
            ("fill out / complete a declaration form", "điền đầy đủ thông tin vào tờ khai"),
            ("submit the declaration form", "nộp tờ khai cho cán bộ"),
            ("sign the declaration form", "ký tên xác nhận vào tờ khai"),
            ("declare currency on the form", "khai báo số tiền mặt mang theo trên tờ khai")
        ],
        "example_vn": "Vui lòng điền tờ khai hải quan trước khi máy bay hạ cánh."
    },
    "duty-free shop": {
        "synonyms": [
            ("tax-free shop", "cửa hàng miễn thuế"),
            ("duty-free store", "tiệm bán hàng miễn thuế")
        ],
        "collocations": [
            ("shop at the duty-free shop", "mua sắm tại cửa hàng miễn thuế"),
            ("duty-free allowance", "hạn mức hàng hóa miễn thuế được phép mang về"),
            ("duty-free goods / items", "mặt hàng miễn thuế (nước hoa, rượu, thuốc lá)"),
            ("airport duty-free zone", "khu mua sắm miễn thuế sân bay")
        ],
        "example_vn": "Tôi đã mua một chai nước hoa tại cửa hàng miễn thuế."
    },
    "exchange booth": {
        "synonyms": [
            ("currency exchange counter", "quầy đổi tiền tệ"),
            ("money changer", "quầy đổi ngoại tệ"),
            ("bureau de change", "quầy đổi tiền (thuật ngữ mượn tiếng Pháp, phổ biến tại châu Âu)")
        ],
        "collocations": [
            ("currency exchange booth", "quầy đổi ngoại tệ sân bay"),
            ("exchange rate at the booth", "tỷ giá quy đổi tại quầy"),
            ("exchange foreign currency", "đổi tiền ngoại tệ sang nội tệ"),
            ("foreign exchange (forex) booth", "quầy giao dịch ngoại hối")
        ],
        "example_vn": "Có một quầy đổi tiền ở ngay gần lối ra."
    },
    "departure lounge": {
        "synonyms": [
            ("waiting lounge", "phòng chờ"),
            ("boarding lounge", "khu phòng chờ lên máy bay"),
            ("gate lounge", "khu vực ghế chờ tại cửa bay"),
            ("concourse lounge", "sảnh chờ hành lang ga")
        ],
        "collocations": [
            ("wait in the departure lounge", "ngồi chờ tại phòng chờ xuất phát"),
            ("business class departure lounge", "phòng chờ hạng thương gia"),
            ("VIP departure lounge", "phòng chờ khách VIP"),
            ("relax in the departure lounge", "nghỉ ngơi thư giãn tại phòng chờ")
        ],
        "example_vn": "Chúng tôi đã chờ ở phòng chờ khởi hành trong suốt hai giờ."
    },

    # ✈️ Flight / Chuyến bay (17 words)
    "flight": {
        "synonyms": [
            ("air trip", "chuyến đi bằng đường hàng không"),
            ("plane journey", "hành trình máy bay"),
            ("aviation travel", "chuyến du hành bằng máy bay")
        ],
        "collocations": [
            ("book / catch / miss a flight", "đặt vé / bắt kịp / lỡ chuyến bay"),
            ("domestic / international flight", "chuyến bay nội địa / quốc tế"),
            ("direct / non-stop flight", "chuyến bay thẳng / bay không dừng"),
            ("long-haul / short-haul flight", "chuyến bay đường dài / chặng ngắn"),
            ("flight number", "số hiệu chuyến bay (vd: VN123, TG456)")
        ],
        "example_vn": "Chuyến bay của tôi sẽ khởi hành lúc 3 giờ chiều."
    },
    "departure": {
        "synonyms": [
            ("takeoff", "sự cất cánh"),
            ("leaving", "sự rời đi"),
            ("outbound flight", "chuyến bay chiều đi")
        ],
        "collocations": [
            ("departure time", "giờ khởi hành"),
            ("departure gate", "cổng khởi hành"),
            ("scheduled departure", "giờ xuất phát theo lịch trình"),
            ("departure board / flight information display", "bảng điện tử thông tin chuyến bay đi"),
            ("departures hall", "sảnh đi của nhà ga")
        ],
        "example_vn": "Giờ khởi hành đã được lùi lại đến 4 giờ chiều."
    },
    "arrival": {
        "synonyms": [
            ("landing", "sự hạ cánh"),
            ("touchdown", "cú tiếp đất của máy bay"),
            ("inbound flight", "chuyến bay chiều đến")
        ],
        "collocations": [
            ("estimated time of arrival (ETA)", "thời gian dự kiến hạ cánh"),
            ("arrival hall / gate", "sảnh đón khách đến / cửa đón khách"),
            ("arrivals board", "bảng thông báo các chuyến bay đến"),
            ("upon arrival", "ngay khi vừa hạ cánh / đến nơi")
        ],
        "example_vn": "Thời gian hạ cánh dự kiến của chúng tôi là 5:30 chiều."
    },
    "destination": {
        "synonyms": [
            ("arrival point", "điểm đến"),
            ("final stop", "điểm dừng chân cuối cùng"),
            ("target location", "địa điểm hướng tới")
        ],
        "collocations": [
            ("final destination", "điểm đến cuối cùng"),
            ("holiday / tourist destination", "điểm đến du lịch nghỉ dưỡng"),
            ("reach one's destination", "tới được nơi cần đến"),
            ("popular destination", "điểm đến nổi tiếng thu hút nhiều khách")
        ],
        "example_vn": "Điểm đến cuối cùng trong hành trình của bạn là đâu?"
    },
    "layover": {
        "synonyms": [
            ("stopover", "điểm dừng chân quá cảnh (thường kéo dài trên 24 giờ)"),
            ("transit", "việc quá cảnh"),
            ("connection", "chặng nối chuyến")
        ],
        "collocations": [
            ("have a layover", "có một chặng quá cảnh"),
            ("3-hour layover", "thời gian quá cảnh kéo dài 3 tiếng"),
            ("overnight layover", "quá cảnh qua đêm"),
            ("layover airport / city", "sân bay / thành phố dừng chân quá cảnh")
        ],
        "example_vn": "Chúng tôi có một chặng quá cảnh kéo dài 3 tiếng tại Singapore."
    },
    "connecting flight": {
        "synonyms": [
            ("transfer flight", "chuyến bay trung chuyển"),
            ("onward flight", "chuyến bay tiếp nối chặng sau")
        ],
        "collocations": [
            ("catch / make a connecting flight", "kịp bắt chuyến bay nối tiếp"),
            ("miss a connecting flight", "bị lỡ chuyến bay nối tiếp"),
            ("tight connecting flight", "khoảng thời gian nối chuyến rất gấp"),
            ("transfer to a connecting flight", "chuyển sang cửa chuyến bay kế tiếp")
        ],
        "example_vn": "Tôi cần bắt chuyến bay nối tiếp để bay sang Tokyo."
    },
    "direct flight": {
        "synonyms": [
            ("non-stop flight", "chuyến bay bay liền mạch không dừng tiếp nhiên liệu"),
            ("through flight", "chuyến bay suốt tuyến (cùng số hiệu máy bay)")
        ],
        "collocations": [
            ("book a direct flight", "đặt vé bay chuyến bay thẳng"),
            ("operate a direct flight", "khai thác đường bay thẳng"),
            ("prefer a direct flight", "chuộng bay thẳng hơn bay quá cảnh")
        ],
        "example_vn": "Đây có phải là chuyến bay bay thẳng tới London không?"
    },
    "takeoff": {
        "synonyms": [
            ("departure", "sự khởi hành bay lên"),
            ("lift-off", "sự nhấc mình khỏi mặt đất"),
            ("ascent", "sự leo độ cao ban đầu")
        ],
        "collocations": [
            ("prepare for takeoff", "chuẩn bị cất cánh"),
            ("cleared for takeoff", "được kiểm soát không lưu cho phép cất cánh"),
            ("smooth takeoff", "pha cất cánh êm ái"),
            ("takeoff runway", "đường băng cất cánh"),
            ("abort takeoff", "hủy thao tác cất cánh khẩn cấp")
        ],
        "example_vn": "Xin vui lòng thắt dây an toàn trước khi máy bay cất cánh."
    },
    "landing": {
        "synonyms": [
            ("touchdown", "cú chạm bánh tiếp đất"),
            ("arrival", "sự hạ cánh đến"),
            ("descent", "quá trình hạ độ cao tiếp cận")
        ],
        "collocations": [
            ("make a landing", "thực hiện hạ cánh"),
            ("smooth / rough landing", "cú tiếp đất êm dịu / dằn xóc"),
            ("emergency landing", "hạ cánh khẩn cấp"),
            ("landing gear", "bộ phận càng đáp bánh xe"),
            ("prepare the cabin for landing", "chuẩn bị khoang hành khách hạ cánh")
        ],
        "example_vn": "Máy bay của chúng tôi đã có một cú hạ cánh vô cùng êm ái."
    },
    "board": {
        "synonyms": [
            ("embark", "bước lên tàu / lên máy bay (trang trọng)"),
            ("get on", "bước lên máy bay (thông dụng)"),
            ("enter the plane", "vào trong máy bay")
        ],
        "collocations": [
            ("board the aircraft / plane", "bước lên máy bay"),
            ("now boarding", "đang cho hành khách lên tàu"),
            ("ready to board", "sẵn sàng lên máy bay"),
            ("priority boarding", "lối ưu tiên lên tàu trước")
        ],
        "example_vn": "Các hành khách ở hàng ghế từ 20 đến 30, xin mời lên máy bay ngay bây giờ."
    },
    "boarding": {
        "synonyms": [
            ("embarkation", "quá trình lên máy bay (thuật ngữ chính thức)"),
            ("getting on board", "việc hành khách lên máy bay")
        ],
        "collocations": [
            ("boarding time", "giờ bắt đầu cho khách lên máy bay"),
            ("boarding gate", "cổng lên tàu bay"),
            ("final boarding call", "thông báo gọi khách lên tàu lần cuối"),
            ("boarding process", "quy trình soát vé lên máy bay")
        ],
        "example_vn": "Việc đón khách lên máy bay sẽ bắt đầu lúc 2:30 chiều."
    },
    "delayed": {
        "synonyms": [
            ("postponed", "bị dời lại thời gian"),
            ("held up", "bị ách lại, kẹt lại"),
            ("deferred", "bị tạm hoãn"),
            ("late", "muộn, trễ")
        ],
        "collocations": [
            ("flight is delayed", "chuyến bay bị hoãn"),
            ("heavily / severely delayed", "bị hoãn rất lâu"),
            ("delayed due to bad weather", "bị trễ do thời tiết xấu"),
            ("delayed by [time]", "bị hoãn mất khoảng thời gian là...")
        ],
        "example_vn": "Chuyến bay đã bị hoãn lại một tiếng đồng hồ."
    },
    "cancelled": {
        "synonyms": [
            ("called off", "bị hủy bỏ"),
            ("grounded", "bị giữ lại trên mặt đất không cho bay"),
            ("scratched", "bị gạch tên khỏi lịch bay")
        ],
        "collocations": [
            ("cancelled flight", "chuyến bay đã bị hủy"),
            ("officially cancelled", "chính thức bị hủy bỏ"),
            ("cancelled due to technical issues", "bị hủy do sự cố kỹ thuật"),
            ("compensation for cancelled flights", "tiền bồi thường chuyến bay bị hủy")
        ],
        "example_vn": "Rất tiếc phải thông báo rằng chuyến bay của bạn đã bị hủy."
    },
    "on time": {
        "synonyms": [
            ("punctual", "đúng giờ chuẩn xác"),
            ("on schedule", "đúng theo lịch trình ấn định"),
            ("promptly", "kịp thời, mau lẹ")
        ],
        "collocations": [
            ("arrive on time", "hạ cánh đúng giờ"),
            ("depart on time", "khởi hành đúng giờ"),
            ("running on time", "hoạt động đúng theo lịch trình"),
            ("on-time performance", "tỷ lệ các chuyến bay đúng giờ của hãng")
        ],
        "example_vn": "Đừng lo lắng, chuyến bay vẫn đang vận hành đúng giờ."
    },
    "turbulence": {
        "synonyms": [
            ("rough air", "luồng không khí động dữ dội"),
            ("air pocket", "túi khí làm máy bay rơi đột ngột"),
            ("atmospheric disturbance", "nhiễu động khí quyển"),
            ("bumpiness", "sự xóc nảy, chòng chành")
        ],
        "collocations": [
            ("experience / hit / encounter turbulence", "gặp phải vùng nhiễu động không khí"),
            ("severe / moderate / light turbulence", "nhiễu động dữ dội / vừa phải / nhẹ"),
            ("clear-air turbulence (CAT)", "nhiễu động trời trong không nhìn thấy được trên radar"),
            ("turbulence warning", "cảnh báo máy bay sắp rung lắc")
        ],
        "example_vn": "Chúng ta đang bay qua một vùng không khí bị nhiễu động nhẹ."
    },
    "announcement": {
        "synonyms": [
            ("notice", "thông báo"),
            ("public notification", "lời thông báo công khai"),
            ("broadcast", "bản tin phát thanh"),
            ("statement", "lời tuyên bố")
        ],
        "collocations": [
            ("make an announcement", "phát thanh đưa ra một thông báo"),
            ("listen to / pay attention to the announcement", "chú ý lắng nghe thông báo"),
            ("public address (PA) announcement", "thông báo qua hệ thống loa phát thanh"),
            ("in-flight announcement", "thông báo trong khoang khách khi đang bay"),
            ("boarding announcement", "thông báo mời hành khách lên tàu")
        ],
        "example_vn": "Xin quý hành khách chú ý lắng nghe thông báo."
    },
    "runway": {
        "synonyms": [
            ("airstrip", "đường băng dã chiến"),
            ("landing strip", "dải đất hạ cánh"),
            ("tarmac", "khu vực mặt sân phủ nhựa đường / bãi đỗ lăn")
        ],
        "collocations": [
            ("take off from the runway", "cất cánh từ đường băng"),
            ("land on the runway", "tiếp đất hạ cánh xuống đường băng"),
            ("active runway", "đường băng đang có máy bay hoạt động"),
            ("runway lights", "hệ thống đèn chiếu sáng đường băng"),
            ("taxi to the runway", "lăn bánh máy bay ra đầu đường băng")
        ],
        "example_vn": "Chiếc máy bay hiện đang đứng chờ trên đường băng."
    },

    # 💺 On the Plane / Trên máy bay (15 words)
    "aisle seat": {
        "synonyms": [
            ("walkway seat", "ghế cạnh lối đi"),
            ("corridor seat", "ghế lối hành lang")
        ],
        "collocations": [
            ("prefer an aisle seat", "thích ngồi ghế cạnh lối đi hơn"),
            ("request an aisle seat", "yêu cầu được xếp ghế cạnh lối đi"),
            ("sit in an aisle seat", "ngồi ở vị trí ghế cạnh lối đi"),
            ("aisle seat advantage", "lợi thế dễ đứng dậy đi vệ sinh hoặc duỗi chân")
        ],
        "example_vn": "Tôi thích ghế cạnh lối đi hơn để có thể duỗi chân thoải mái."
    },
    "window seat": {
        "synonyms": [
            ("window-side seat", "chỗ ngồi cạnh cửa sổ")
        ],
        "collocations": [
            ("request / book a window seat", "yêu cầu / đặt ghế cạnh cửa sổ"),
            ("enjoy the view from a window seat", "ngắm mây trời phong cảnh từ ghế cửa sổ"),
            ("pull down / close the window shade", "kéo tấm chắn sáng cửa sổ xuống")
        ],
        "example_vn": "Làm ơn cho tôi một ghế ngồi cạnh cửa sổ được không?"
    },
    "middle seat": {
        "synonyms": [
            ("center seat", "ghế trung tâm giữa hai ghế"),
            ("between seat", "ghế kẹp giữa")
        ],
        "collocations": [
            ("stuck in the middle seat", "bị kẹt ngồi ở ghế giữa"),
            ("avoid the middle seat", "tránh đặt trúng ghế giữa"),
            ("assigned a middle seat", "bị xếp vào vị trí ghế giữa"),
            ("armrest battle for the middle seat", "quy ước nhường hai chỗ để tay cho người ngồi ghế giữa")
        ],
        "example_vn": "Hầu như không ai thích phải ngồi ở ghế giữa cả."
    },
    "overhead bin": {
        "synonyms": [
            ("overhead compartment", "ngăn hành lý trên đầu"),
            ("overhead locker", "tủ hành lý trên trần (chuẩn Anh-Anh)"),
            ("luggage bin", "khoang chứa đồ trên cao")
        ],
        "collocations": [
            ("put / stow in the overhead bin", "cất hành lý vào hộc trên đầu"),
            ("close / latch the overhead bin", "đóng chốt nắp ngăn hành lý lại"),
            ("full / crowded overhead bin", "hộc hành lý đã chật kín"),
            ("space in the overhead bin", "khoảng trống trong ngăn hành lý")
        ],
        "example_vn": "Vui lòng đặt túi xách của bạn vào ngăn hành lý phía trên đầu."
    },
    "tray table": {
        "synonyms": [
            ("fold-down table", "bàn gập hạ xuống"),
            ("seatback table", "bàn gắn sau lưng ghế"),
            ("meal tray", "khay để đồ ăn")
        ],
        "collocations": [
            ("fold up / stow the tray table", "gập gọn bàn ăn lại"),
            ("lower / put down the tray table", "hạ bàn ăn xuống phía trước"),
            ("lock the tray table in upright position", "chốt bàn ăn ở vị trí dựng thẳng đứng")
        ],
        "example_vn": "Vui lòng gập bàn ăn lại để máy bay chuẩn bị hạ cánh."
    },
    "seatbelt": {
        "synonyms": [
            ("safety belt", "đai an toàn"),
            ("lap belt", "dây an toàn thắt ngang hông"),
            ("restraint", "thiết bị giữ an toàn")
        ],
        "collocations": [
            ("fasten your seatbelt", "cài thắt dây an toàn"),
            ("unfasten / unbuckle the seatbelt", "tháo mở chốt dây an toàn"),
            ("keep your seatbelt fastened", "duy trì thắt dây an toàn trong suốt chuyến bay"),
            ("fasten seatbelt sign", "đèn hiệu báo hiệu yêu cầu thắt dây an toàn")
        ],
        "example_vn": "Xin quý hành khách vui lòng cài dây an toàn."
    },
    "fasten": {
        "synonyms": [
            ("buckle", "cài chốt khóa"),
            ("secure", "buộc / siết chặt an toàn"),
            ("latch", "gài chốt"),
            ("tie", "buộc lại")
        ],
        "collocations": [
            ("fasten securely", "thắt gài một cách chắc chắn"),
            ("fasten seatbelts", "thắt dây an toàn"),
            ("ensure seatbelts are fastened", "đảm bảo dây an toàn đã được cài đúng cách")
        ],
        "example_vn": "Vui lòng thắt dây an toàn trong suốt quá trình cất cánh."
    },
    "emergency exit": {
        "synonyms": [
            ("fire exit", "cửa thoát hiểm khi hỏa hoạn"),
            ("evacuation door", "cửa sơ tán hành khách"),
            ("escape hatch", "cửa thoát hiểm trên thân máy bay")
        ],
        "collocations": [
            ("emergency exit row", "hàng ghế ngay sát cửa thoát hiểm (chỗ để chân rộng)"),
            ("locate the nearest emergency exit", "xác định vị trí cửa thoát hiểm gần nhất"),
            ("never block the emergency exit", "không bao giờ để hành lý chắn lối thoát hiểm"),
            ("emergency exit slide", "máng trượt sơ tán khẩn cấp")
        ],
        "example_vn": "Lối thoát hiểm nằm ngay phía trên cánh máy bay."
    },
    "life vest": {
        "synonyms": [
            ("life jacket", "áo phao cứu sinh"),
            ("life preserver", "phao cứu sinh cá nhân"),
            ("personal flotation device (PFD)", "thiết bị nổi cứu sinh cá nhân")
        ],
        "collocations": [
            ("inflate the life vest", "giật chốt làm phồng áo phao (khi đã ra khỏi cửa máy bay)"),
            ("put on the life vest", "mặc áo phao vào người"),
            ("life vest located under your seat", "áo phao được trang bị ngay dưới gầm ghế ngồi"),
            ("whistle and emergency light on the vest", "còi và đèn cứu nạn gắn trên áo phao")
        ],
        "example_vn": "Áo phao cứu sinh được đặt ngay dưới gầm ghế của bạn."
    },
    "blanket": {
        "synonyms": [
            ("throw", "chăn đắp mỏng"),
            ("fleece", "chăn nỉ giữ nhiệt"),
            ("quilt", "chăn bông"),
            ("coverlet", "tấm phủ giữ ấm")
        ],
        "collocations": [
            ("ask for an extra blanket", "xin thêm một chiếc chăn"),
            ("airline blanket", "chăn phục vụ của hãng hàng không"),
            ("warm blanket", "chiếc chăn ấm áp"),
            ("wrap up in a blanket", "quấn chăn quanh người để ngủ")
        ],
        "example_vn": "Cho tôi xin một chiếc chăn được không? Trong khoang hơi lạnh."
    },
    "pillow": {
        "synonyms": [
            ("cushion", "đệm tựa / gối đệm"),
            ("headrest", "gối tựa đầu trên ghế máy bay"),
            ("travel pillow", "gối cổ du lịch chữ U")
        ],
        "collocations": [
            ("ask for an extra pillow", "xin thêm một chiếc gối"),
            ("travel neck pillow", "gối kê cổ hình chữ U"),
            ("rest your head on the pillow", "tựa đầu vào gối nghỉ ngơi")
        ],
        "example_vn": "Tôi có thể xin thêm một chiếc gối nữa được không?"
    },
    "headphones": {
        "synonyms": [
            ("earphones", "tai nghe nhét tai nhỏ gọn"),
            ("headset", "bộ tai nghe trùm đầu kèm micro"),
            ("in-ear monitors", "tai nghe in-ear")
        ],
        "collocations": [
            ("noise-cancelling headphones", "tai nghe chống ồn chủ động cao cấp"),
            ("airline-provided headphones", "tai nghe do hãng hàng không phát"),
            ("plug in headphones", "cắm tai nghe vào giắc cắm trên ghế"),
            ("headphone jack", "lỗ cắm tai nghe trên thành ghế máy bay")
        ],
        "example_vn": "Tai nghe nằm trong túi ghế phía trước mặt bạn."
    },
    "beverage": {
        "synonyms": [
            ("drink", "thức uống thông thường"),
            ("refreshment", "đồ uống giải khát thanh mát"),
            ("potable", "nước uống được (trang trọng)")
        ],
        "collocations": [
            ("complimentary beverage", "đồ uống phục vụ miễn phí trên máy bay"),
            ("hot / cold beverage", "thức uống nóng (trà, cà phê) / lạnh (nước ngọt, nước ép)"),
            ("alcoholic / non-alcoholic beverage", "đồ uống có cồn (rượu, bia) / không cồn"),
            ("beverage service / cart", "xe đẩy phục vụ đồ uống của tiếp viên")
        ],
        "example_vn": "Bạn có muốn dùng một món đồ uống nào không?"
    },
    "meal": {
        "synonyms": [
            ("in-flight meal", "suất ăn trên máy bay"),
            ("dish", "món ăn"),
            ("repast", "bữa ăn (văn phong trang trọng)")
        ],
        "collocations": [
            ("hot meal", "suất ăn nóng hổi"),
            ("meal service", "khung giờ phục vụ bữa ăn chính"),
            ("special meal (vegetarian, halal, kosher, gluten-free)", "suất ăn đặc biệt (chay, đạo Hồi, Kosher, không gluten)"),
            ("pre-order a meal", "đặt trước suất ăn khi mua vé")
        ],
        "example_vn": "Bữa ăn sẽ được tiếp viên phục vụ trong khoảng 30 phút nữa."
    },
    "seat pocket": {
        "synonyms": [
            ("seatback pocket", "túi chứa đồ sau lưng ghế trước"),
            ("seat pouch", "túi vải gắn sau ghế")
        ],
        "collocations": [
            ("safety briefing card in the seat pocket", "thẻ hướng dẫn an toàn trong túi ghế"),
            ("in-flight magazine in the seat pocket", "tạp chí của hãng bay trong túi ghế"),
            ("stow items in the seat pocket", "nhét vật dụng cá nhân vào túi ghế"),
            ("check the seat pocket before disembarking", "kiểm tra túi ghế trước khi xuống máy bay tránh bỏ quên đồ")
        ],
        "example_vn": "Thẻ hướng dẫn an toàn bay nằm ngay trong túi ghế phía trước."
    },

    # 🚕 Transportation / Di chuyển (4 words)
    "taxi": {
        "synonyms": [
            ("cab / taxicab", "xe taxi (chuẩn tiếng Anh-Mỹ)"),
            ("ride-hailing car", "xe công nghệ đặt qua app (Grab, Uber, Gojek)")
        ],
        "collocations": [
            ("take / catch / hail a taxi", "bắt / gọi một chiếc taxi"),
            ("taxi stand / taxi rank", "bến đỗ đón taxi tại sân bay"),
            ("taxi meter", "đồng hồ tính cước taxi"),
            ("taxi fare", "giá cước taxi"),
            ("licensed airport taxi", "taxi sân bay chính hãng, được cấp phép")
        ],
        "example_vn": "Chúng ta hãy bắt một chiếc taxi để về khách sạn nhé."
    },
    "subway": {
        "synonyms": [
            ("underground", "hệ thống tàu điện ngầm (chuẩn Anh-Anh)"),
            ("tube", "tàu điện ngầm (tiếng lóng quen thuộc tại London)"),
            ("metro", "tàu điện ngầm (phổ biến tại châu Âu và thế giới)"),
            ("MRT / BTS", "hệ thống đường sắt đô thị (tại Singapore, Bangkok)")
        ],
        "collocations": [
            ("take / ride the subway", "đi lại bằng tàu điện ngầm"),
            ("subway station", "nhà ga tàu điện ngầm"),
            ("subway line / route", "tuyến đường tàu điện ngầm"),
            ("subway ticket / smart card", "vé tàu / thẻ từ đi tàu điện ngầm"),
            ("airport express subway link", "tuyến tàu điện ngầm tốc hành kết nối sân bay")
        ],
        "example_vn": "Đi tàu điện ngầm thì tiết kiệm chi phí hơn nhiều so với đi taxi."
    },
    "shuttle bus": {
        "synonyms": [
            ("airport shuttle", "xe buýt trung chuyển sân bay"),
            ("transfer bus", "xe buýt đưa đón chặng ngắn"),
            ("feeder bus", "xe buýt trung chuyển kết nối ga tàu/sân bay")
        ],
        "collocations": [
            ("free / complimentary shuttle bus", "xe buýt đưa đón miễn phí"),
            ("hotel shuttle bus", "xe buýt đón tiễn khách của khách sạn"),
            ("airport terminal shuttle bus", "xe buýt di chuyển giữa các nhà ga sân bay"),
            ("catch / board the shuttle bus", "đón và lên xe shuttle bus"),
            ("shuttle bus schedule / timetable", "lịch trình chạy của xe shuttle bus")
        ],
        "example_vn": "Có một tuyến xe buýt đưa đón miễn phí chạy thẳng vào trung tâm thành phố."
    },
    "rental car": {
        "synonyms": [
            ("car rental", "dịch vụ thuê ô tô"),
            ("hire car", "xe thuê (chuẩn Anh-Anh)"),
            ("rented vehicle", "phương tiện đi thuê")
        ],
        "collocations": [
            ("reserve / book a rental car", "đặt trước một chiếc xe tự lái"),
            ("rental car counter / agency", "quầy đại lý cho thuê xe tại sảnh sân bay"),
            ("pick up / drop off a rental car", "nhận xe / trả xe thuê"),
            ("rental car insurance", "bảo hiểm xe thuê"),
            ("rental car agreement / contract", "hợp đồng thuê xe")
        ],
        "example_vn": "Tôi đã đặt trước một chiếc xe ô tô tự lái tại sân bay."
    }
}

# XML Helper functions
def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=140, right=140):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def set_cell_border(cell, **kwargs):
    """
    kwargs can be top, bottom, left, right.
    val="single", sz="4", color="CBD5E1"
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}/>')
    for edge in ('top', 'left', 'bottom', 'right'):
        edge_data = kwargs.get(edge)
        if edge_data:
            b = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="{edge_data.get("val","single")}" w:sz="{edge_data.get("sz","4")}" w:space="0" w:color="{edge_data.get("color","CBD5E1")}"/>')
            tcBorders.append(b)
        else:
            b = parse_xml(f'<w:{edge} {nsdecls("w")} w:val="none"/>')
            tcBorders.append(b)
    tcPr.append(tcBorders)

def set_table_borders(table, color="CBD5E1", sz="4"):
    tblPr = table._tbl.tblPr
    borders = parse_xml(f'<w:tblBorders {nsdecls("w")}><w:top w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:bottom w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:left w:val="none"/><w:right w:val="none"/><w:insideH w:val="single" w:sz="{sz}" w:space="0" w:color="{color}"/><w:insideV w:val="none"/></w:tblBorders>')
    tblPr.append(borders)

def set_row_header(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_row_cant_split(row):
    trPr = row._tr.get_or_add_trPr()
    trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))

def set_col_widths(table, col_widths_dxa):
    for row in table.rows:
        for idx, width in enumerate(col_widths_dxa):
            if idx < len(row.cells):
                tcPr = row.cells[idx]._tc.get_or_add_tcPr()
                # Remove existing tcW if any
                for w in tcPr.findall(qn('w:tcW')):
                    tcPr.remove(w)
                tcW = parse_xml(f'<w:tcW {nsdecls("w")} w:w="{width}" w:type="dxa"/>')
                tcPr.append(tcW)


def parse_airport_md(file_path):
    with open(file_path, "r", encoding="utf-8") as f:
        content = f.read()

    # Split into vocab part and conversation part
    parts = content.split("# 2. Conversations / Hội thoại")
    vocab_text = parts[0]
    conv_text = "# 2. Conversations / Hội thoại" + parts[1]

    # Parse vocab categories
    cat_blocks = re.findall(r'## (.*?)\n\n\| Từ vựng \| IPA \| Nghĩa \| Ví dụ \|\n\|[-|]+\|\n(.*?)(?=\n## |\n---|\Z)', vocab_text, re.DOTALL)
    categories = []
    for cat_name, table_body in cat_blocks:
        cat_name_clean = cat_name.strip()
        words = []
        for line in table_body.strip().split("\n"):
            line = line.strip()
            if not line.startswith("|"):
                continue
            cells = [c.strip() for c in line.split("|")[1:-1]]
            if len(cells) >= 4:
                raw_word = cells[0]
                ipa = cells[1]
                meaning = cells[2]
                example = cells[3].replace("**", "")
                words.append({
                    "raw_word": raw_word,
                    "ipa": ipa,
                    "meaning": meaning,
                    "example": example
                })
        categories.append((cat_name_clean, words))

    # Parse conversations
    sections = []
    major_groups = re.findall(r'## (2\.\d+\..*?)(?=\n## 2\.|\Z)', conv_text, re.DOTALL)
    for group_block in major_groups:
        group_title = group_block.strip().split("\n")[0].strip()
        conv_blocks = re.findall(r'### (Hội thoại \d+:.*?)(?=\n### Hội thoại |\Z)', group_block, re.DOTALL)
        dialogues = []
        for cb in conv_blocks:
            lines = cb.strip().split("\n")
            title_line = lines[0].strip()
            # Context
            context_match = re.search(r'>\s*\*Tình huống:\s*(.*?)\*', cb)
            context = context_match.group(1).strip() if context_match else ""
            
            # English dialogue
            en_match = re.search(r'#### 🇬🇧 Tiếng Anh\s*\n(.*?)(?=#### 🇻🇳 Tiếng Việt|\Z)', cb, re.DOTALL)
            en_lines = []
            if en_match:
                for l in en_match.group(1).strip().split("\n"):
                    l_str = l.strip()
                    if l_str.startswith(">") and len(l_str) > 2:
                        cleaned = l_str.lstrip("> ").strip()
                        if cleaned:
                            en_lines.append(cleaned)
                            
            # Vietnamese dialogue
            vn_match = re.search(r'#### 🇻🇳 Tiếng Việt\s*\n(.*?)(?=\n---|\Z)', cb, re.DOTALL)
            vn_lines = []
            if vn_match:
                for l in vn_match.group(1).strip().split("\n"):
                    l_str = l.strip()
                    if l_str.startswith(">") and len(l_str) > 2:
                        cleaned = l_str.lstrip("> ").strip()
                        if cleaned:
                            vn_lines.append(cleaned)

            dialogues.append({
                "title": title_line,
                "context": context,
                "en_lines": en_lines,
                "vn_lines": vn_lines
            })
        sections.append((group_title, dialogues))

    return categories, sections


def build_word_document():
    md_path = "D:\\Desktop\\English in Conservation\\airport.md"
    docx_path = "D:\\Desktop\\English in Conservation\\airport.docx"

    print("Parsing airport.md...")
    categories, sections = parse_airport_md(md_path)
    print(f"Parsed {len(categories)} categories and {len(sections)} conversation groups.")

    doc = docx.Document()

    # Page Setup: A4, Margins: Top=0.7", Bottom=0.7", Left=0.6", Right=0.6"
    for section in doc.sections:
        section.page_width = Inches(8.27)   # A4 Width
        section.page_height = Inches(11.69) # A4 Height
        section.top_margin = Inches(0.7)
        section.bottom_margin = Inches(0.7)
        section.left_margin = Inches(0.6)
        section.right_margin = Inches(0.6)
        
        # Header / Footer
        header = section.header
        hp = header.paragraphs[0]
        hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        hrun = hp.add_run("English in Conversation: At the Airport ✈️ | Full Study Guide")
        hrun.font.name = "Segoe UI"
        hrun.font.size = Pt(8.5)
        hrun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

        footer = section.footer
        fp = footer.paragraphs[0]
        fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        frun = fp.add_run("Tiếng Anh Giao Tiếp Sân Bay • Từ vựng mở rộng, Từ đồng nghĩa, Collocations & Hội thoại thực tế")
        frun.font.name = "Segoe UI"
        frun.font.size = Pt(8)
        frun.font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

    # Base Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Segoe UI'
    normal_style.font.size = Pt(10)
    normal_style.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    # ==================== DOCUMENT COVER / TITLE BLOCK ====================
    # Title box table
    banner_table = doc.add_table(rows=1, cols=1)
    banner_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    banner_cell = banner_table.rows[0].cells[0]
    set_cell_background(banner_cell, "1E3A8A") # Navy Blue
    set_cell_margins(banner_cell, top=240, bottom=240, left=260, right=260)
    set_cell_border(banner_cell, 
                    top={"val": "single", "sz": "12", "color": "0D9488"},
                    bottom={"val": "single", "sz": "12", "color": "0D9488"})

    p_banner_tag = banner_cell.paragraphs[0]
    p_banner_tag.paragraph_format.space_before = Pt(0)
    p_banner_tag.paragraph_format.space_after = Pt(4)
    run_tag = p_banner_tag.add_run("✈️ ENGLISH IN CONVERSATION • PRACTICAL TRAVEL GUIDE")
    run_tag.font.name = "Segoe UI"
    run_tag.font.size = Pt(9.5)
    run_tag.font.bold = True
    run_tag.font.color.rgb = RGBColor(0x38, 0xBD, 0xF8) # Sky light

    p_banner_title = banner_cell.add_paragraph()
    p_banner_title.paragraph_format.space_before = Pt(2)
    p_banner_title.paragraph_format.space_after = Pt(6)
    run_title = p_banner_title.add_run("AT THE AIRPORT — TẠI SÂN BAY")
    run_title.font.name = "Segoe UI"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

    p_banner_sub = banner_cell.add_paragraph()
    p_banner_sub.paragraph_format.space_before = Pt(0)
    p_banner_sub.paragraph_format.space_after = Pt(4)
    run_sub = p_banner_sub.add_run("Cẩm Nang Từ Vựng Mở Rộng (Synonyms & Collocations) & 12 Tình Huống Hội Thoại Song Ngữ Thực Tế")
    run_sub.font.name = "Segoe UI"
    run_sub.font.size = Pt(11.5)
    run_sub.font.color.rgb = RGBColor(0xE2, 0xE8, 0xF0)

    # Empty space after banner
    p_space = doc.add_paragraph()
    p_space.paragraph_format.space_before = Pt(6)
    p_space.paragraph_format.space_after = Pt(6)

    # Overview Callout
    intro_table = doc.add_table(rows=1, cols=1)
    intro_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    intro_cell = intro_table.rows[0].cells[0]
    set_cell_background(intro_cell, "F0F9FF") # very light blue
    set_cell_margins(intro_cell, top=140, bottom=140, left=180, right=180)
    set_cell_border(intro_cell, 
                    left={"val": "single", "sz": "18", "color": "0284C7"},
                    top=None, right=None, bottom=None)
    
    p_intro = intro_cell.paragraphs[0]
    p_intro.paragraph_format.space_before = Pt(0)
    p_intro.paragraph_format.space_after = Pt(2)
    r_intro_bold = p_intro.add_run("📖 Hướng dẫn sử dụng tài liệu: \n")
    r_intro_bold.bold = True
    r_intro_bold.font.color.rgb = RGBColor(0x03, 0x69, 0xA1)
    r_intro_text = p_intro.add_run(
        "• Tài liệu này được biên soạn toàn diện dựa trên giáo trình airport.md, nâng cấp với hệ thống Từ đồng nghĩa (Synonyms) và Cụm từ thường gặp (Collocations) cho 100% từ vựng.\n"
        "• Phần 1 gồm 66 từ vựng cốt lõi chia theo 5 chủ đề: Con người, Sân bay, Chuyến bay, Trên máy bay, Phương tiện di chuyển.\n"
        "• Phần 2 gồm 12 đoạn hội thoại thực tế được trình bày song ngữ song song (Side-by-Side) tương ứng từ khâu làm thủ tục (Check-in), kiểm tra an ninh, trên máy bay đến nhập cảnh và lấy hành lý."
    )
    r_intro_text.font.size = Pt(9.5)
    r_intro_text.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

    # ==================== PART 1: ENRICHED VOCABULARY ====================
    p_h1 = doc.add_paragraph()
    p_h1.paragraph_format.space_before = Pt(22)
    p_h1.paragraph_format.space_after = Pt(8)
    r_h1 = p_h1.add_run("1. VOCABULARY / TỪ VỰNG CHUYÊN SÂU & MỞ RỘNG")
    r_h1.font.name = "Segoe UI"
    r_h1.font.size = Pt(16)
    r_h1.font.bold = True
    r_h1.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    # Column widths for 5-column table (Total width: ~7.07 inches = 10180 dxa)
    # Col 0 (Từ vựng & IPA): 1850 dxa (~1.28 in)
    # Col 1 (Nghĩa): 1400 dxa (~0.97 in)
    # Col 2 (Synonyms): 1850 dxa (~1.28 in)
    # Col 3 (Collocations): 2580 dxa (~1.79 in)
    # Col 4 (Ví dụ & Dịch): 2500 dxa (~1.74 in)
    col_widths = [1850, 1400, 1850, 2580, 2500]

    for cat_idx, (cat_name, words) in enumerate(categories, start=1):
        p_cat = doc.add_paragraph()
        p_cat.paragraph_format.space_before = Pt(14)
        p_cat.paragraph_format.space_after = Pt(6)
        r_cat = p_cat.add_run(f"1.{cat_idx}. {cat_name} ({len(words)} từ vựng)")
        r_cat.font.name = "Segoe UI"
        r_cat.font.size = Pt(12.5)
        r_cat.font.bold = True
        r_cat.font.color.rgb = RGBColor(0x02, 0x84, 0xC7)

        # Create Table
        table = doc.add_table(rows=1, cols=5)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders(table, color="CBD5E1", sz="4")

        # Header Row
        hdr_row = table.rows[0]
        set_row_header(hdr_row)
        headers = [
            "Từ vựng & IPA",
            "Nghĩa tiếng Việt",
            "Từ đồng nghĩa (Synonyms)",
            "Cụm từ thông dụng (Collocations)",
            "Ví dụ & Dịch câu"
        ]
        for col_idx, h_text in enumerate(headers):
            c = hdr_row.cells[col_idx]
            set_cell_background(c, "1E3A8A")
            set_cell_margins(c, top=120, bottom=120, left=100, right=100)
            p = c.paragraphs[0]
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            r = p.add_run(h_text)
            r.font.name = "Segoe UI"
            r.font.size = Pt(9.5)
            r.font.bold = True
            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

        # Data Rows
        for w_idx, w in enumerate(words):
            raw_w = w["raw_word"]
            # Extract plain word key for dictionary lookup
            base_key = re.sub(r'\s*\([a-z.]+\)\s*', '', raw_w).strip().lower()
            enrich_info = ENRICHED_DATA.get(base_key, {})

            row = table.add_row()
            set_row_cant_split(row)
            bg_color = "F8FAFC" if w_idx % 2 == 1 else "FFFFFF"

            for c in row.cells:
                set_cell_background(c, bg_color)
                set_cell_margins(c, top=90, bottom=90, left=90, right=90)

            # Col 0: Từ vựng & IPA
            c0 = row.cells[0]
            p0 = c0.paragraphs[0]
            p0.paragraph_format.space_before = Pt(0)
            p0.paragraph_format.space_after = Pt(2)
            r0_word = p0.add_run(raw_w)
            r0_word.bold = True
            r0_word.font.size = Pt(9.5)
            r0_word.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)

            p0_ipa = c0.add_paragraph()
            p0_ipa.paragraph_format.space_before = Pt(0)
            p0_ipa.paragraph_format.space_after = Pt(0)
            r0_ipa = p0_ipa.add_run(w["ipa"])
            r0_ipa.font.size = Pt(8.5)
            r0_ipa.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)

            # Col 1: Nghĩa tiếng Việt
            c1 = row.cells[1]
            p1 = c1.paragraphs[0]
            p1.paragraph_format.space_before = Pt(0)
            p1.paragraph_format.space_after = Pt(0)
            r1 = p1.add_run(w["meaning"])
            r1.font.size = Pt(9)
            r1.font.bold = True
            r1.font.color.rgb = RGBColor(0x1E, 0x29, 0x3B)

            # Col 2: Synonyms
            c2 = row.cells[2]
            syns = enrich_info.get("synonyms", [])
            if syns:
                for s_i, (syn_word, syn_desc) in enumerate(syns):
                    p_syn = c2.paragraphs[0] if s_i == 0 else c2.add_paragraph()
                    p_syn.paragraph_format.space_before = Pt(0)
                    p_syn.paragraph_format.space_after = Pt(2)
                    r_sb = p_syn.add_run(f"• {syn_word}")
                    r_sb.bold = True
                    r_sb.font.size = Pt(8.5)
                    r_sb.font.color.rgb = RGBColor(0x03, 0x69, 0xA1) # blue
                    if syn_desc:
                        r_sd = p_syn.add_run(f": {syn_desc}")
                        r_sd.font.size = Pt(8)
                        r_sd.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            else:
                p_syn = c2.paragraphs[0]
                p_syn.add_run("—").font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

            # Col 3: Collocations
            c3 = row.cells[3]
            collocs = enrich_info.get("collocations", [])
            if collocs:
                for c_i, (col_term, col_meaning) in enumerate(collocs):
                    p_col = c3.paragraphs[0] if c_i == 0 else c3.add_paragraph()
                    p_col.paragraph_format.space_before = Pt(0)
                    p_col.paragraph_format.space_after = Pt(2)
                    r_cb = p_col.add_run(f"• {col_term}")
                    r_cb.bold = True
                    r_cb.font.size = Pt(8.5)
                    r_cb.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) # teal
                    if col_meaning:
                        r_cm = p_col.add_run(f"\n  ({col_meaning})")
                        r_cm.font.size = Pt(7.8)
                        r_cm.font.color.rgb = RGBColor(0x47, 0x55, 0x69)
            else:
                p_col = c3.paragraphs[0]
                p_col.add_run("—").font.color.rgb = RGBColor(0x94, 0xA3, 0xB8)

            # Col 4: Ví dụ & Dịch
            c4 = row.cells[4]
            p4_en = c4.paragraphs[0]
            p4_en.paragraph_format.space_before = Pt(0)
            p4_en.paragraph_format.space_after = Pt(2)
            r4_en = p4_en.add_run(w["example"])
            r4_en.font.size = Pt(8.5)
            r4_en.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

            vn_trans = enrich_info.get("example_vn", "")
            if vn_trans:
                p4_vn = c4.add_paragraph()
                p4_vn.paragraph_format.space_before = Pt(1)
                p4_vn.paragraph_format.space_after = Pt(0)
                r4_vn = p4_vn.add_run(f"→ {vn_trans}")
                r4_vn.font.size = Pt(8)
                r4_vn.font.italic = True
                r4_vn.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)

        set_col_widths(table, col_widths)

    # ==================== PART 2: CONVERSATIONS ====================
    doc.add_page_break()

    p_h2 = doc.add_paragraph()
    p_h2.paragraph_format.space_before = Pt(16)
    p_h2.paragraph_format.space_after = Pt(6)
    r_h2 = p_h2.add_run("2. PRACTICAL CONVERSATIONS / 12 ĐOẠN HỘI THOẠI THỰC HÀNH")
    r_h2.font.name = "Segoe UI"
    r_h2.font.size = Pt(16)
    r_h2.font.bold = True
    r_h2.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

    p_conv_intro = doc.add_paragraph()
    p_conv_intro.paragraph_format.space_before = Pt(0)
    p_conv_intro.paragraph_format.space_after = Pt(12)
    r_ci = p_conv_intro.add_run("Dưới đây là 12 tình huống hội thoại giao tiếp thực tế xuyên suốt hành trình bay. Mỗi đoạn hội thoại được trình bày dạng bảng song ngữ đối xứng Anh - Việt, phân tách rõ vai người nói (Nhân viên / Hành khách) giúp việc luyện đọc và đối chiếu trở nên trực quan, sinh động nhất.")
    r_ci.font.size = Pt(9.5)
    r_ci.font.color.rgb = RGBColor(0x47, 0x55, 0x69)

    # Column widths for 2-column conversation table: 5090 dxa + 5090 dxa = 10180 dxa
    conv_col_widths = [5090, 5090]

    for group_title, dialogues in sections:
        p_grp = doc.add_paragraph()
        p_grp.paragraph_format.space_before = Pt(16)
        p_grp.paragraph_format.space_after = Pt(8)
        r_grp = p_grp.add_run(group_title)
        r_grp.font.name = "Segoe UI"
        r_grp.font.size = Pt(13)
        r_grp.font.bold = True
        r_grp.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)

        for d in dialogues:
            p_dt = doc.add_paragraph()
            p_dt.paragraph_format.space_before = Pt(12)
            p_dt.paragraph_format.space_after = Pt(4)
            r_dt = p_dt.add_run(d["title"])
            r_dt.font.name = "Segoe UI"
            r_dt.font.size = Pt(11.5)
            r_dt.font.bold = True
            r_dt.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)

            # Context Box
            if d["context"]:
                c_tbl = doc.add_table(rows=1, cols=1)
                c_tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                c_cell = c_tbl.rows[0].cells[0]
                set_cell_background(c_cell, "F8FAFC")
                set_cell_margins(c_cell, top=80, bottom=80, left=140, right=140)
                set_cell_border(c_cell, left={"val": "single", "sz": "12", "color": "3B82F6"})
                
                cp = c_cell.paragraphs[0]
                cp.paragraph_format.space_before = Pt(0)
                cp.paragraph_format.space_after = Pt(0)
                r_ctx_tag = cp.add_run("📌 Tình huống: ")
                r_ctx_tag.bold = True
                r_ctx_tag.font.size = Pt(9)
                r_ctx_tag.font.color.rgb = RGBColor(0x1E, 0x40, 0xAF)
                r_ctx_val = cp.add_run(d["context"])
                r_ctx_val.font.size = Pt(9)
                r_ctx_val.font.italic = True
                r_ctx_val.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

                p_sp = doc.add_paragraph()
                p_sp.paragraph_format.space_before = Pt(2)
                p_sp.paragraph_format.space_after = Pt(2)

            # Dialogue Table
            d_table = doc.add_table(rows=1, cols=2)
            d_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            set_table_borders(d_table, color="CBD5E1", sz="4")

            # Table Header
            d_hdr = d_table.rows[0]
            set_row_header(d_hdr)
            
            c_en_h = d_hdr.cells[0]
            set_cell_background(c_en_h, "334155") # Slate dark
            set_cell_margins(c_en_h, top=100, bottom=100, left=120, right=120)
            p_enh = c_en_h.paragraphs[0]
            r_enh = p_enh.add_run("🇬🇧 Tiếng Anh (English Dialogue)")
            r_enh.bold = True
            r_enh.font.size = Pt(9.5)
            r_enh.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            c_vn_h = d_hdr.cells[1]
            set_cell_background(c_vn_h, "334155")
            set_cell_margins(c_vn_h, top=100, bottom=100, left=120, right=120)
            p_vnh = c_vn_h.paragraphs[0]
            r_vnh = p_vnh.add_run("🇻🇳 Bản Dịch Tiếng Việt")
            r_vnh.bold = True
            r_vnh.font.size = Pt(9.5)
            r_vnh.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

            # Rows for each speaker turn
            num_turns = max(len(d["en_lines"]), len(d["vn_lines"]))
            for t_idx in range(num_turns):
                t_row = d_table.add_row()
                set_row_cant_split(t_row)
                row_bg = "FFFFFF" if t_idx % 2 == 0 else "F8FAFC"

                # English cell
                c_en = t_row.cells[0]
                set_cell_background(c_en, row_bg)
                set_cell_margins(c_en, top=80, bottom=80, left=100, right=100)
                p_en = c_en.paragraphs[0]
                p_en.paragraph_format.space_before = Pt(0)
                p_en.paragraph_format.space_after = Pt(0)

                en_text = d["en_lines"][t_idx] if t_idx < len(d["en_lines"]) else ""
                if en_text:
                    # Parse speaker badge (e.g. 🛎️ **Staff:** ...)
                    # Check if begins with emoji/speaker
                    spk_match = re.match(r'^(🛎️|\👤)?\s*\*\*(.*?)\*\*:?\s*(.*)$', en_text)
                    if spk_match:
                        icon = spk_match.group(1) or ""
                        spk_name = spk_match.group(2).rstrip(':').strip()
                        spk_utterance = spk_match.group(3)

                        r_badge = p_en.add_run(f"{icon} {spk_name}: ".strip() + " ")
                        r_badge.bold = True
                        r_badge.font.size = Pt(9)
                        if "Staff" in spk_name or "🛎️" in icon:
                            r_badge.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Navy
                        else:
                            r_badge.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E) # Teal

                        r_utt = p_en.add_run(spk_utterance)
                        r_utt.font.size = Pt(9)
                        r_utt.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)
                    else:
                        r_plain = p_en.add_run(en_text.replace("**", ""))
                        r_plain.font.size = Pt(9)
                        r_plain.font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

                # Vietnamese cell
                c_vn = t_row.cells[1]
                set_cell_background(c_vn, row_bg)
                set_cell_margins(c_vn, top=80, bottom=80, left=100, right=100)
                p_vn = c_vn.paragraphs[0]
                p_vn.paragraph_format.space_before = Pt(0)
                p_vn.paragraph_format.space_after = Pt(0)

                vn_text = d["vn_lines"][t_idx] if t_idx < len(d["vn_lines"]) else ""
                if vn_text:
                    spk_match_vn = re.match(r'^(🛎️|\👤)?\s*\*\*(.*?)\*\*:?\s*(.*)$', vn_text)
                    if spk_match_vn:
                        icon = spk_match_vn.group(1) or ""
                        spk_name = spk_match_vn.group(2).rstrip(':').strip()
                        spk_utterance = spk_match_vn.group(3)

                        r_badge = p_vn.add_run(f"{icon} {spk_name}: ".strip() + " ")
                        r_badge.bold = True
                        r_badge.font.size = Pt(9)
                        if "Nhân viên" in spk_name or "🛎️" in icon:
                            r_badge.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
                        else:
                            r_badge.font.color.rgb = RGBColor(0x0F, 0x76, 0x6E)

                        r_utt = p_vn.add_run(spk_utterance)
                        r_utt.font.size = Pt(9)
                        r_utt.font.color.rgb = RGBColor(0x33, 0x41, 0x55)
                    else:
                        r_plain = p_vn.add_run(vn_text.replace("**", ""))
                        r_plain.font.size = Pt(9)
                        r_plain.font.color.rgb = RGBColor(0x33, 0x41, 0x55)

            set_col_widths(d_table, conv_col_widths)

    # ==================== SAVE DOCUMENT ====================
    doc.save(docx_path)
    print(f"Successfully created and saved {docx_path}!")

if __name__ == "__main__":
    build_word_document()
