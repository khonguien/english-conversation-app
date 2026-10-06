import os
import sys
import shutil
from PIL import Image
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

prs = Presentation("Airport_backup_before_images.pptx")
print(f"Loaded presentation with {len(prs.slides)} slides.")

TARGET_BOX_LEFT = 7753800
TARGET_BOX_TOP = 1800000
TARGET_BOX_WIDTH = 3800000
TARGET_BOX_HEIGHT = 4300000
BOX_ASPECT = TARGET_BOX_WIDTH / TARGET_BOX_HEIGHT

images = os.listdir("downloaded_images")

inserted_count = 0
for slide_num in range(25, 68):
    matching = [f for f in images if f.startswith(f"slide_{slide_num}_")]
    if not matching:
        print(f"Slide {slide_num}: No image file found!")
        continue
    
    img_path = os.path.join("downloaded_images", matching[0])
    slide_idx = slide_num - 1
    slide = prs.slides[slide_idx]
    
    try:
        with Image.open(img_path) as img:
            w_px, h_px = img.size
        
        aspect = w_px / h_px
        if aspect > BOX_ASPECT:
            final_width = TARGET_BOX_WIDTH
            final_height = int(TARGET_BOX_WIDTH / aspect)
            final_left = TARGET_BOX_LEFT
            final_top = TARGET_BOX_TOP + (TARGET_BOX_HEIGHT - final_height) // 2
        else:
            final_height = TARGET_BOX_HEIGHT
            final_width = int(TARGET_BOX_HEIGHT * aspect)
            final_left = TARGET_BOX_LEFT + (TARGET_BOX_WIDTH - final_width) // 2
            final_top = TARGET_BOX_TOP
        
        slide.shapes.add_picture(img_path, final_left, final_top, final_width, final_height)
        inserted_count += 1
        print(f"Slide {slide_num}: Inserted {matching[0]} ({w_px}x{h_px})")
    except Exception as e:
        print(f"Slide {slide_num} error: {e}")

# Save to Airport_with_images.pptx
out_path = "Airport_with_images.pptx"
prs.save(out_path)
print(f"\nSuccessfully inserted {inserted_count}/43 images and saved to {out_path}!")

# Try to overwrite Airport.pptx if unlocked
try:
    shutil.copyfile(out_path, "Airport.pptx")
    print("Successfully updated Airport.pptx as well!")
except PermissionError:
    print("Notice: Airport.pptx is currently open in PowerPoint, so could not overwrite it directly. Airport_with_images.pptx was created successfully!")
