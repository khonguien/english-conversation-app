import os
import sys
import json
import shutil
import urllib.request
import urllib.parse
from PIL import Image
from pptx import Presentation

sys.stdout.reconfigure(encoding='utf-8')

# Backup presentation before adding images
shutil.copyfile("Airport.pptx", "Airport_backup_before_images.pptx")
print("Backed up to Airport_backup_before_images.pptx")

headers = {'User-Agent': 'AirportVocabDownloader/1.0 (educational_project@gmail.com)'}
os.makedirs("downloaded_images", exist_ok=True)

queries = {
    25: ("baggage claim", ["baggage claim airport", "baggage claim area Heathrow", "baggage claim terminal"]),
    26: ("conveyor belt", ["baggage carousel airport", "baggage conveyor belt airport"]),
    27: ("customs", ["customs border airport", "customs inspection airport counter"]),
    28: ("immigration", ["passport control immigration airport", "airport immigration checkpoint"]),
    29: ("declaration form", ["customs declaration form", "passenger customs form"]),
    30: ("duty-free shop", ["duty free shop airport", "Narita airport duty free", "duty free airport"]),
    31: ("exchange booth", ["currency exchange airport booth", "money exchange airport", "currency exchange counter"]),
    32: ("flight", ["commercial airplane flying in sky", "airliner in flight blue sky"]),
    33: ("departure", ["airport departures board", "departure board airport"]),
    34: ("arrival", ["airport arrivals board", "arrival board airport"]),
    35: ("destination", ["flight information display system airport destinations", "airport flight departure board screen"]),
    36: ("layover", ["airport transit lounge passengers", "passengers waiting airport gate terminal"]),
    37: ("connecting flight", ["flight connections airport sign", "connecting flights transfer airport", "flight connections"]),
    38: ("direct flight", ["passenger airliner in flight", "commercial aircraft in flight cruising"]),
    39: ("takeoff", ["airplane taking off runway", "airliner takeoff runway", "Boeing takeoff runway"]),
    40: ("landing", ["airplane landing runway", "aircraft landing runway", "airliner landing"]),
    41: ("board", ["passengers boarding aircraft", "boarding airplane stairs", "passengers boarding plane"]),
    42: ("boarding", ["airport boarding gate passengers", "boarding gate airport", "gate boarding"]),
    43: ("delayed", ["flight delayed screen airport", "delayed flight board screen"]),
    44: ("cancelled", ["flight cancellations airport screen", "flight cancelled board airport", "cancelled flight"]),
    45: ("on time", ["flight status board on time", "flight information display on time"]),
    46: ("turbulence", ["airplane wing storm clouds", "airplane flying clouds turbulence"]),
    47: ("announcement", ["flight attendant public address announcement", "airport public address announcement speaker"]),
    48: ("runway", ["airport runway markings tarmac", "airport runway asphalt"]),
    49: ("aisle seat", ["airplane cabin aisle seat passenger", "airplane aisle seat"]),
    50: ("window seat", ["airplane window seat view wing", "airplane window view clouds"]),
    51: ("middle seat", ["Lufthansa economy cabin row of seats", "airplane passenger cabin three seats economy"]),
    52: ("overhead bin", ["aircraft overhead bin luggage", "airplane overhead locker passenger"]),
    53: ("tray table", ["airplane tray table folded passenger", "in-flight tray table"]),
    54: ("seat pocket", ["airplane seat pocket safety card", "aircraft seat pocket magazine"]),
    55: ("blanket", ["airline blanket pillow economy seat", "passenger airplane blanket"]),
    56: ("pillow", ["airline pillow airplane seat", "airplane passenger pillow"]),
    57: ("headphones", ["airline in flight entertainment headphones", "KLM headphones airplane"]),
    58: ("seatbelt", ["fastened seatbelt airplane passenger", "airplane seat belt buckle"]),
    59: ("fasten", ["fasten seat belt sign airplane cabin", "fasten seatbelt sign Boeing"]),
    60: ("emergency exit", ["Airbus emergency exit door", "aircraft emergency exit door interior"]),
    61: ("life vest", ["life jacket aircraft", "aviation life vest under seat", "aircraft life vest safety"]),
    62: ("beverage", ["in-flight beverage drink service", "flight attendant drink cart"]),
    63: ("meal", ["in-flight meal tray airline food", "airline meal tray"]),
    64: ("taxi", ["airport taxi terminal queue", "taxi stand airport terminal"]),
    65: ("subway", ["airport metro station train", "airport subway terminal"]),
    66: ("shuttle bus", ["airport shuttle bus terminal", "airport bus terminal passengers"]),
    67: ("rental car", ["Enterprise Rent-a-Car Rental Counter airport", "airport car rental counter desk"])
}

def download_image(queries_list, save_path):
    for q in queries_list:
        url = f"https://commons.wikimedia.org/w/api.php?action=query&generator=search&gsrsearch={urllib.parse.quote(q)}&gsrnamespace=6&gsrlimit=5&prop=imageinfo&iiprop=url|mime|size&iiurlwidth=800&format=json"
        req = urllib.request.Request(url, headers=headers)
        try:
            with urllib.request.urlopen(req, timeout=10) as resp:
                data = json.loads(resp.read().decode('utf-8'))
                pages = data.get('query', {}).get('pages', {})
                for pid, pinfo in pages.items():
                    title = pinfo.get('title', '')
                    ii = pinfo.get('imageinfo', [{}])[0]
                    mime = ii.get('mime', '')
                    thumb = ii.get('thumburl', '')
                    if mime in ['image/jpeg', 'image/png'] and thumb:
                        img_req = urllib.request.Request(thumb, headers=headers)
                        with urllib.request.urlopen(img_req, timeout=15) as img_resp:
                            with open(save_path, 'wb') as f:
                                f.write(img_resp.read())
                        with Image.open(save_path) as img:
                            w, h = img.size
                        return True, title, w, h
        except Exception as e:
            pass
    return False, None, 0, 0

prs = Presentation("Airport.pptx")

# Target bounding box on the right side
TARGET_BOX_LEFT = 7753800
TARGET_BOX_TOP = 1800000
TARGET_BOX_WIDTH = 3800000
TARGET_BOX_HEIGHT = 4300000
BOX_ASPECT = TARGET_BOX_WIDTH / TARGET_BOX_HEIGHT

success_count = 0

for slide_num, (word, q_list) in queries.items():
    slide_idx = slide_num - 1
    slide = prs.slides[slide_idx]
    
    clean_word = word.replace(' ', '_').replace('-', '_')
    save_path = os.path.join("downloaded_images", f"slide_{slide_num}_{clean_word}.jpg")
    
    ok, title, w_px, h_px = download_image(q_list, save_path)
    if ok and w_px > 0 and h_px > 0:
        # Calculate fit
        img_aspect = w_px / h_px
        if img_aspect > BOX_ASPECT:
            final_width = TARGET_BOX_WIDTH
            final_height = int(TARGET_BOX_WIDTH / img_aspect)
            final_left = TARGET_BOX_LEFT
            final_top = TARGET_BOX_TOP + (TARGET_BOX_HEIGHT - final_height) // 2
        else:
            final_height = TARGET_BOX_HEIGHT
            final_width = int(TARGET_BOX_HEIGHT * img_aspect)
            final_left = TARGET_BOX_LEFT + (TARGET_BOX_WIDTH - final_width) // 2
            final_top = TARGET_BOX_TOP
        
        # Add picture to slide
        slide.shapes.add_picture(save_path, final_left, final_top, final_width, final_height)
        success_count += 1
        safe_title = title.encode('ascii', 'replace').decode('ascii') if title else ''
        print(f"Slide {slide_num} ({word}): Inserted successfully -> {safe_title[:45]}")
    else:
        print(f"Slide {slide_num} ({word}): FAILED to download image")

prs.save("Airport.pptx")
print(f"\nFinished! Added images to {success_count}/{len(queries)} slides.")
